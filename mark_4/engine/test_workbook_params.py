"""Workbook integration checks without loading prices or running a backtest."""
from dataclasses import asdict
import os
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

from config import read_strategy_params
from runtime import load_shared
from workbook_controls import preserve_runner_controls
from workbook_params import preserve_parameter_sheet

WORKBOOK = Path(os.environ.get('MARK4_TEST_WORKBOOK', Path(__file__).resolve().parents[2] / 'Backtest_Model_Runner.xlsm'))


@unittest.skipUnless(WORKBOOK.exists(), 'Mark 4 workbook not supplied')
class ParameterWorkbook(unittest.TestCase):
    def test_current_settings_and_layout(self):
        wb = load_workbook(WORKBOOK, data_only=False)
        try:
            params = read_strategy_params(wb)
            self.assertEqual(params.pivot_scales, (2, 3, 5))
            self.assertEqual(params.historical_pivot_lookback_bars, 1024)
            self.assertEqual(params.entry_retracement, 0.618)
            self.assertEqual(params.min_imbalance_pips, 4)
            self.assertFalse(params.only_first_filled_trade_per_day)
            s = wb['Strategy_Params']
            self.assertIsNotNone(s.freeze_panes)
            self.assertFalse(params.use_consolidation_buy_setup)
            self.assertFalse(params.use_consolidation_obstacles)
            self.assertTrue(params.use_pivot_obstacles)
            self.assertTrue(params.use_active_price_cluster_stop_filter)
            self.assertTrue(params.reject_22_only_setups)
            self.assertEqual(params.pivot_obstacle_expiry_months, 24)
            self.assertEqual(params.price_cluster_expiry_months, 24)
            self.assertTrue(params.use_repeated_visit_pivot_exception)
            self.assertTrue(params.use_trendline_obstacles)
            self.assertTrue(params.use_trendline_cluster_score_filter)
            self.assertEqual((params.trendline_trade_side_oldest_local_min_score,
                params.trendline_trade_side_latest_local_min_score,
                params.trendline_opposite_side_latest_trailing_min_score), (70, 65, 50))
            self.assertEqual(params.repeated_visit_minimum_visits, 3)
            self.assertEqual(params.repeated_visit_membership_tolerance_pips, 4)
            self.assertEqual(params.repeated_visit_departure_pips, 5)
            self.assertFalse(bool(s.row_dimensions[92].hidden))
            self.assertEqual(s['C32'].number_format, '0.0%')
            self.assertEqual(s['C72'].value, 0.7)
            self.assertEqual(s['E37'].value, '15m signal bars')
        finally:
            wb.close()

    def test_validation_has_one_correct_rule_per_input(self):
        wb = load_workbook(WORKBOOK)
        try:
            s = wb['Strategy_Params']
            for row in range(4, s.max_row+1):
                kind = s[f'D{row}'].value
                rules = [rule for rule in s.data_validations.dataValidation if f'C{row}' in rule]
                if kind in ('Bool', 'Int', 'Float'):
                    self.assertEqual(len(rules), 1, row)
                    self.assertEqual(rules[0].type, {'Bool':'list', 'Int':'whole', 'Float':'decimal'}[kind], row)
                    self.assertTrue(rules[0].showErrorMessage, row)
            self.assertEqual(next(r for r in s.data_validations.dataValidation if 'C16' in r).formula1, '"Yes,No"')
            self.assertEqual(next(r for r in s.data_validations.dataValidation if 'C38' in r).formula1, '"Yes"')
        finally:
            wb.close()

    def test_shared_writer_roundtrip_keeps_layout_and_current_values(self):
        writer = load_shared('excel_writer')
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / 'roundtrip.xlsm'
            wb = load_workbook(WORKBOOK, keep_vba=True)
            before = asdict(read_strategy_params(wb))
            original_hidden = wb['Strategy_Params'].row_dimensions[8].hidden
            original_freeze = wb['Strategy_Params'].freeze_panes
            # Deliberately vary one destination value to test that presentation
            # restoration does not copy stale input parameters over run values.
            wb['Strategy_Params']['C30'] = 6
            writer.ensure_strategy_params_sheet(wb)
            wb.save(output)
            wb.close()
            preserve_runner_controls(WORKBOOK, output)
            preserve_parameter_sheet(WORKBOOK, output)
            result = load_workbook(output)
            try:
                expected = before | {'min_imbalance_pips': 6.0}
                self.assertEqual(asdict(read_strategy_params(result)), expected)
                s = result['Strategy_Params']
                self.assertEqual(s['A1'].value, 'Mark 4 Strategy Parameters')
                self.assertEqual(s['E37'].value, '15m signal bars')
                self.assertEqual(s.freeze_panes, original_freeze)
                self.assertEqual(s.row_dimensions[8].hidden, original_hidden)
                rules = [r for r in s.data_validations.dataValidation if 'C30' in r]
                self.assertEqual(len(rules), 1)
                self.assertEqual(rules[0].type, 'decimal')
                self.assertEqual(s['C32'].number_format, '0.0%')
            finally:
                result.close()
            with ZipFile(WORKBOOK) as src, ZipFile(output) as dst:
                self.assertEqual(src.read('xl/vbaProject.bin'), dst.read('xl/vbaProject.bin'))
                for name in src.namelist():
                    if name.startswith(('xl/ctrlProps/', 'xl/drawings/')):
                        self.assertEqual(src.read(name), dst.read(name), name)


if __name__ == '__main__':
    unittest.main()
