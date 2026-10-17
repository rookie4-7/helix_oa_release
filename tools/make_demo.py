"""Generate synthetic candles for a plumbing demonstration, never market evidence."""
import csv
from datetime import datetime, timedelta
import math
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "demo_data"
HEADERS = ["time", *[f"{kind}_{field}" for kind in ("mid", "bid", "ask")
                     for field in ("open", "high", "low", "close")], "volume"]


def write_prices(path, rows, history=False):
    book = Workbook(write_only=True)
    if history:
        for year in (2020, 2021):
            book.create_sheet(str(year)).append(HEADERS)
    sheet = book.create_sheet("2022")
    sheet.append(HEADERS)
    for stamp, o, h, l, c, volume in rows:
        values = [stamp]
        for offset in (0, -0.00005, 0.00005):
            values.extend(round(p + offset, 5) for p in (o, h, l, c))
        sheet.append([*values, volume])
    book.save(path)


def main():
    if TARGET.exists():
        raise SystemExit("demo_data already exists; preserve or rename it before generating another demo.")
    TARGET.mkdir()
    (TARGET / "GBP_USD_M1_MBA").mkdir()
    start = datetime(2022, 1, 3)
    rows = []
    previous = 1.35
    for minute in range(3 * 24 * 60):
        close = round(1.35 + 0.0015 * math.sin(minute / 110) + 0.0002 * math.sin(minute / 7), 5)
        rows.append((start + timedelta(minutes=minute), previous,
                     max(previous, close) + 0.00003, min(previous, close) - 0.00003,
                     close, 20 + minute % 15))
        previous = close
    coarse = []
    for i in range(0, len(rows), 15):
        chunk = rows[i:i + 15]
        coarse.append((chunk[0][0], chunk[0][1], max(r[2] for r in chunk),
                       min(r[3] for r in chunk), chunk[-1][4], sum(r[5] for r in chunk)))
    write_prices(TARGET / "GBP_USD_M1_MBA/GBP_USD_M1_MBA_2022.xlsx", rows)
    write_prices(TARGET / "GBP_USD_M15_MBA_2015_2025_by_year.xlsx", coarse, history=True)
    with (TARGET / "news_calendar.csv").open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["datetime", "timezone", "currency", "impact", "event"])
        writer.writerow(["2022-01-05 19:00:00", "UTC", "USD", "High", "SYNTHETIC demo event - not real news"])
    print("Synthetic demo created. Not real prices, news, or performance evidence.")


if __name__ == "__main__":
    main()
