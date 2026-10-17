from __future__ import annotations

from dataclasses import fields
from dataclasses import replace
from typing import Any

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

try:
    from .models import SelectedWindow
    from .models import SampleType
    from .models import StrategyParams
except ImportError:
    from models import SelectedWindow
    from models import SampleType
    from models import StrategyParams


REQUIRED_SHEETS = {
    "Config",
    "Data_Catalog",
    "Data_Quality_Report",
    "Signal_Model",
    "Trade_Log",
    "Walk_Forward",
    "Run_Log",
    "IS_Analysis",
    "OOS_Analysis",
    "WF_Window_Analysis",
    "Dashboard",
    "Notes_Change_Log",
}


class WorkbookConfigError(Exception):
    """Raised when the Excel model config cannot be read safely."""


def require_sheets(workbook: Workbook) -> None:
    missing = sorted(REQUIRED_SHEETS - set(workbook.sheetnames))
    if missing:
        raise WorkbookConfigError(f"Workbook is missing required sheets: {', '.join(missing)}")


def read_config(sheet: Worksheet) -> dict[str, Any]:
    config: dict[str, Any] = {}
    for row_idx in range(1, sheet.max_row + 1):
        key = sheet.cell(row=row_idx, column=1).value
        if isinstance(key, str) and key.strip():
            config[key.strip()] = sheet.cell(row=row_idx, column=2).value
    return config


def table_headers(sheet: Worksheet, header_row: int) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col_idx in range(1, sheet.max_column + 1):
        value = sheet.cell(row=header_row, column=col_idx).value
        if isinstance(value, str) and value.strip():
            headers[value.strip()] = col_idx
    return headers


def read_table(
    sheet: Worksheet,
    header_row: int,
    first_data_row: int,
    last_data_row: int | None = None,
) -> list[dict[str, Any]]:
    headers = table_headers(sheet, header_row)
    if not headers:
        return []

    max_row = min(last_data_row or sheet.max_row, sheet.max_row)
    rows: list[dict[str, Any]] = []
    for row_idx in range(first_data_row, max_row + 1):
        row = {
            name: sheet.cell(row=row_idx, column=col_idx).value
            for name, col_idx in headers.items()
        }
        if any(value not in (None, "") for value in row.values()):
            row["_excel_row"] = row_idx
            rows.append(row)
    return rows


def selected_window(workbook: Workbook, window_id: str) -> SelectedWindow:
    windows = read_table(workbook["Walk_Forward"], header_row=3, first_data_row=4, last_data_row=11)
    for row in windows:
        if row.get("Window_ID") == window_id:
            return SelectedWindow(
                window_id=str(row["Window_ID"]),
                is_start=row.get("IS_Start"),
                is_end=row.get("IS_End"),
                oos_start=row.get("OOS_Start"),
                oos_end=row.get("OOS_End"),
                oos_year=row.get("OOS_Year"),
                target_rule_version=str(row.get("Target_Rule_Version") or ""),
            )
    raise WorkbookConfigError(f"Selected window {window_id!r} was not found in Walk_Forward.")


def samples_to_run(sample: str) -> list[SampleType]:
    if sample == "Both":
        return ["IS", "OOS"]
    if sample == "IS":
        return ["IS"]
    if sample == "OOS":
        return ["OOS"]
    raise WorkbookConfigError("Run_Sample must be IS, OOS, or Both.")


def read_rule_config(workbook: Workbook, config: dict[str, Any]) -> dict[str, Any]:
    rules = read_table(workbook["Signal_Model"], header_row=4, first_data_row=5, last_data_row=36)
    active_rules = [
        {key: value for key, value in row.items() if key != "_excel_row"}
        for row in rules
        if str(row.get("Enabled") or "").strip().lower() == "yes"
    ]
    return {
        "workbook_version": config.get("Workbook Version"),
        "strategy_name": config.get("Strategy Name"),
        "rule_version": config.get("Rule_Version"),
        "rule_set_id": config.get("Rule_Set_ID"),
        "active_rules": active_rules,
    }


def read_strategy_params(workbook: Workbook, base: StrategyParams | None = None) -> StrategyParams:
    """Read editable Strategy_Params rows into the Python strategy dataclass."""
    params = base or StrategyParams()
    if "Strategy_Params" not in workbook.sheetnames:
        return params

    rows = read_table(workbook["Strategy_Params"], header_row=3, first_data_row=4)
    valid_names = {field.name for field in fields(StrategyParams)}
    updates: dict[str, Any] = {}
    for row in rows:
        name = str(row.get("Parameter") or "").strip()
        if not name:
            continue
        if name not in valid_names:
            raise WorkbookConfigError(f"Unknown Strategy_Params parameter: {name!r}")
        raw_value = row.get("Value")
        current_value = getattr(params, name)
        updates[name] = parse_strategy_param_value(name, raw_value, current_value)

    return replace(params, **updates)


def parse_strategy_param_value(name: str, raw_value: Any, default_value: Any) -> Any:
    if raw_value in (None, ""):
        return default_value

    if isinstance(default_value, bool):
        return parse_bool(raw_value, name)
    if isinstance(default_value, int) and not isinstance(default_value, bool):
        try:
            return int(float(raw_value))
        except (TypeError, ValueError) as exc:
            raise WorkbookConfigError(f"Strategy_Params {name!r} must be an integer.") from exc
    if isinstance(default_value, float):
        try:
            return float(raw_value)
        except (TypeError, ValueError) as exc:
            raise WorkbookConfigError(f"Strategy_Params {name!r} must be numeric.") from exc
    if isinstance(default_value, tuple):
        if isinstance(raw_value, str):
            return tuple(part.strip() for part in raw_value.split(",") if part.strip())
        if isinstance(raw_value, (list, tuple)):
            return tuple(str(part).strip() for part in raw_value if str(part).strip())
        return (str(raw_value).strip(),)
    return str(raw_value).strip()


def parse_bool(raw_value: Any, name: str) -> bool:
    if isinstance(raw_value, bool):
        return raw_value
    if isinstance(raw_value, (int, float)):
        return bool(raw_value)
    if isinstance(raw_value, str):
        normalized = raw_value.strip().lower()
        if normalized in {"yes", "y", "true", "t", "1", "on", "enabled"}:
            return True
        if normalized in {"no", "n", "false", "f", "0", "off", "disabled"}:
            return False
    raise WorkbookConfigError(f"Strategy_Params {name!r} must be Yes/No or True/False.")
