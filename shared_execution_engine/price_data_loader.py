from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

try:
    from .data_quality import missing_column_stats
    from .data_quality import missing_file_stats
    from .data_quality import missing_required_columns
    from .data_quality import OPTIONAL_MBA_COLUMNS
    from .data_quality import quality_status
    from .data_quality import quick_price_stats
    from .data_quality import REQUIRED_MBA_COLUMNS
    from .data_quality import scan_price_sheet
    from .data_quality import valid_date_bounds
    from .models import PriceFile
    from .models import PriceFileStats
except ImportError:
    from data_quality import missing_column_stats
    from data_quality import missing_file_stats
    from data_quality import missing_required_columns
    from data_quality import OPTIONAL_MBA_COLUMNS
    from data_quality import quality_status
    from data_quality import quick_price_stats
    from data_quality import REQUIRED_MBA_COLUMNS
    from data_quality import scan_price_sheet
    from data_quality import valid_date_bounds
    from models import PriceFile
    from models import PriceFileStats

SYMBOL = "GBP_USD"
M1_DIR_NAME = "GBP_USD_M1_MBA"
M15_FILE_NAME = "GBP_USD_M15_MBA_2015_2025_by_year.xlsx"


class PriceDataError(Exception):
    """Raised when local price data files are missing or malformed."""


def to_jsonable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Path):
        return str(value)
    return value


def years_between(start: Any, end: Any) -> list[int]:
    start_year = _coerce_year(start)
    end_year = _coerce_year(end)
    if start_year is None or end_year is None:
        return list(range(2015, 2026))
    return list(range(start_year, end_year + 1))


def discover_price_files(project_root: Path, years: list[int] | None = None) -> list[PriceFile]:
    root = Path(project_root).expanduser().resolve()
    years_to_include = years or list(range(2015, 2026))
    files: list[PriceFile] = []

    m1_dir = root / M1_DIR_NAME
    for year in years_to_include:
        files.append(
            PriceFile(
                symbol=SYMBOL,
                timeframe="1m",
                year=year,
                path=m1_dir / f"GBP_USD_M1_MBA_{year}.xlsx",
                sheet_name=str(year),
                source_type="yearly_excel",
            )
        )

    m15_path = root / M15_FILE_NAME
    for year in years_to_include:
        files.append(
            PriceFile(
                symbol=SYMBOL,
                timeframe="15m",
                year=year,
                path=m15_path,
                sheet_name=str(year),
                source_type="multi_sheet_excel",
            )
        )

    return files


def inspect_price_file(price_file: PriceFile, scan_rows: bool = False, max_scan_rows: int | None = None) -> PriceFileStats:
    if not price_file.path.exists():
        return missing_file_stats(price_file)

    workbook = load_workbook(price_file.path, read_only=True, data_only=True)
    try:
        if price_file.sheet_name not in workbook.sheetnames:
            raise PriceDataError(f"{price_file.path} is missing sheet {price_file.sheet_name!r}.")

        sheet = workbook[price_file.sheet_name]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in header_row]
        missing = missing_required_columns(headers)
        row_count = max(0, (sheet.max_row or 1) - 1)

        if missing:
            return missing_column_stats(price_file, headers, row_count, missing)

        if scan_rows:
            return scan_price_sheet(price_file, sheet, headers, max_scan_rows)

        time_column = headers.index("time") + 1
        start_time = _first_time(sheet, time_column)
        end_time = _last_time(sheet, time_column)
        return quick_price_stats(price_file, headers, row_count, start_time, end_time)
    finally:
        workbook.close()


def iter_mba_rows(
    price_file: PriceFile,
    start: Any | None = None,
    end: Any | None = None,
    max_rows: int | None = None,
) -> Iterator[dict[str, Any]]:
    workbook = load_workbook(price_file.path, read_only=True, data_only=True)
    try:
        if price_file.sheet_name not in workbook.sheetnames:
            raise PriceDataError(f"{price_file.path} is missing sheet {price_file.sheet_name!r}.")

        sheet = workbook[price_file.sheet_name]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in header_row]
        missing = [column for column in REQUIRED_MBA_COLUMNS if column not in headers]
        if missing:
            raise PriceDataError(
                f"{price_file.path} sheet {price_file.sheet_name} is missing columns: {', '.join(missing)}"
            )

        yielded = 0
        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values))
            timestamp = row.get("time")
            if start is not None and timestamp is not None and timestamp < start:
                continue
            if end is not None and timestamp is not None and timestamp > end:
                continue
            yield row
            yielded += 1
            if max_rows is not None and yielded >= max_rows:
                return
    finally:
        workbook.close()


def inspect_price_files(
    price_files: list[PriceFile],
    scan_rows: bool = False,
    max_scan_rows: int | None = None,
) -> list[PriceFileStats]:
    return [
        inspect_price_file(price_file, scan_rows=scan_rows, max_scan_rows=max_scan_rows)
        for price_file in price_files
    ]


def stats_to_manifest_item(stats: PriceFileStats, project_root: Path) -> dict[str, Any]:
    price_file = stats.price_file
    file_size = price_file.path.stat().st_size if stats.exists and price_file.path.exists() else None
    diagnostics = stats.diagnostics
    core_status = quality_status(stats)
    # These diagnostics are warnings only, never new execution eligibility rules.
    report_status = "Review" if core_status == "Validated" and diagnostics.get("warnings") else core_status
    return {
        "data_id": f"{price_file.symbol}_{price_file.timeframe}_{price_file.year}",
        "symbol": price_file.symbol,
        "timeframe": price_file.timeframe,
        "year": price_file.year,
        "data_type": "MBA",
        "file_path": _relative_or_absolute(price_file.path, project_root),
        "resolved_path": str(price_file.path),
        "file_format": "xlsx",
        "sheet_name": price_file.sheet_name,
        "source_type": price_file.source_type,
        "start_date": stats.start_time,
        "end_date": stats.end_time,
        "rows": stats.rows,
        "timezone": "UTC",
        "timezone_checked": bool(diagnostics.get("timezone_check")),
        "data_version": f"{price_file.symbol}_{price_file.timeframe}_{price_file.year}",
        "checksum": None,
        "status": core_status,
        "report_status": report_status,
        "diagnostics": diagnostics,
        "exists": stats.exists,
        "size_bytes": file_size,
        "headers": stats.headers,
        "missing_required_columns": stats.missing_required_columns,
        "duplicate_timestamps": stats.duplicate_timestamps,
        "out_of_order_timestamps": stats.out_of_order_timestamps,
        "bad_bid_ask_rows": stats.bad_bid_ask_rows,
        "negative_spread_rows": stats.negative_spread_rows,
        "rows_checked": stats.rows_checked,
        "validation_performed": True,
        "date_range_ok": valid_date_bounds(stats) and not stats.out_of_order_timestamps,
        "date_range_scope": "Recorded timestamp bounds within the named year sheet; not calendar completeness",
        "row_check_scope": "Timestamp duplicates/order and bid/ask/mid close quotes",
    }


def _first_time(sheet: Any, time_column: int = 1) -> Any:
    return sheet.cell(row=2, column=time_column).value if (sheet.max_row or 0) >= 2 else None


def _last_time(sheet: Any, time_column: int = 1) -> Any:
    return sheet.cell(row=sheet.max_row, column=time_column).value if (sheet.max_row or 0) >= 2 else None


def _coerce_year(value: Any) -> int | None:
    if isinstance(value, datetime):
        return value.year
    if isinstance(value, date):
        return value.year
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str) and len(value) >= 4 and value[:4].isdigit():
        return int(value[:4])
    return None


def _relative_or_absolute(path: Path, root: Path) -> str:
    try:
        return str(path.resolve().relative_to(root.resolve()))
    except ValueError:
        return str(path)
