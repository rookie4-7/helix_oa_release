from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

def dispatch_versioned_engine() -> None:
    """Existing workbook buttons select an isolated engine without modifying VBA."""
    if "--workbook" not in sys.argv:
        return
    from openpyxl import load_workbook
    path = Path(sys.argv[sys.argv.index("--workbook") + 1]).expanduser().resolve()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        config = {row[0]: row[1] for row in workbook["Config"].iter_rows(values_only=True)
                  if len(row) >= 2 and row[0]}
    finally:
        workbook.close()
    profile = config.get("Engine_Profile")
    versions = {"MARK_2_CURRENT_ENGINE": "mark_2", "MARK_3_CURRENT_ENGINE": "mark_3",
                "MARK_4_CURRENT_ENGINE": "mark_4"}
    if profile in versions:
        import os
        target = Path(__file__).resolve().parents[1] / versions[profile] / "engine/excel_runner.py"
        if not target.is_file():
            raise SystemExit("Selected engine is missing: " + str(target))
        os.execv(sys.executable, [sys.executable, "-B", str(target), *sys.argv[1:]])
    if profile not in (None, "", "CURRENT_ENGINE") or str(config.get("Rule_Version", "")).startswith(("Mark2", "Mark3", "Mark4")):
        raise SystemExit("Unknown or missing engine profile; refusing to run the wrong engine.")


if __name__ == "__main__":
    dispatch_versioned_engine()

import backtest_engine


DEFAULT_PROJECT_ROOT = Path(__file__).resolve().parents[1]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Excel-facing wrapper for the GBP/USD backtest engine.",
    )
    parser.add_argument("--workbook", required=True, help="Open runner workbook path.")
    parser.add_argument(
        "--mode",
        choices=["dry-run", "write"],
        default="dry-run",
        help="Dry run validates/calculates only; write creates a timestamped output workbook.",
    )
    parser.add_argument(
        "--project-root",
        default=str(DEFAULT_PROJECT_ROOT),
        help="Repository root containing the engine, price data, and news calendar.",
    )
    parser.add_argument(
        "--output-root",
        default=None,
        help="Folder for timestamped write-mode workbook outputs.",
    )
    parser.add_argument(
        "--output-workbook",
        default=None,
        help="Exact write-mode output workbook path. Optional; defaults to timestamped output-root path.",
    )
    parser.add_argument("--price-data-root", help="Optional external price-data directory; read only.")
    parser.add_argument("--news-data-root", help="Optional external news-data directory; read only.")
    parser.add_argument("--archive-root", help="Optional new archive directory; must not already exist.")
    parser.add_argument(
        "--activate",
        action="store_true",
        default=True,
        help="Mark generated run rows Active_For_Analysis = Yes in write mode.",
    )
    parser.add_argument(
        "--official-oos",
        action="store_true",
        help="Mark generated OOS rows Official_OOS_Run = Yes in write mode.",
    )
    parser.add_argument(
        "--disable-news-filter",
        action="store_true",
        help="Run without news-caused imbalance invalidation.",
    )
    parser.add_argument(
        "--disable-news-forced-close",
        action="store_true",
        help="Run without pre-news forced closes on open trades.",
    )
    return parser.parse_args()


def to_jsonable(value: Any) -> Any:
    return backtest_engine.to_jsonable(value)


def output_workbook_path(workbook_path: Path, output_root: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    suffix = workbook_path.suffix if workbook_path.suffix.lower() in {".xlsm", ".xlsx"} else ".xlsm"
    base_stem = workbook_path.stem.split("_run_", 1)[0]
    return output_root / f"{base_stem}_run_{timestamp}{suffix}"


def build_engine_args(args: argparse.Namespace) -> argparse.Namespace:
    workbook_path = Path(args.workbook).expanduser().resolve()
    project_root = Path(args.project_root).expanduser().resolve()
    output_root = Path(args.output_root or project_root / "excel_runner_runs").expanduser().resolve()
    write = args.mode == "write"
    output_root.mkdir(parents=True, exist_ok=True)
    output_path = (Path(args.output_workbook).expanduser().resolve() if args.output_workbook
                   else output_workbook_path(workbook_path, output_root)) if write else None
    if output_path and output_path.exists():
        raise backtest_engine.EngineError(f"Refusing to overwrite an existing output: {output_path}")
    archive_root = getattr(args, "archive_root", None)
    if write and not archive_root:
        archive_root = output_root / "archive" / output_path.stem
    if write and archive_root and Path(archive_root).exists():
        raise backtest_engine.EngineError(f"Choose a new archive directory: {archive_root}")

    return argparse.Namespace(
        workbook=str(workbook_path),
        output_workbook=str(output_path) if output_path else None,
        project_root=str(project_root),
        archive_root=str(Path(archive_root).expanduser().resolve()) if archive_root else None,
        window_id=None,
        sample=None,
        write=write,
        activate=bool(args.activate) if write else False,
        official_oos=bool(args.official_oos) if write else False,
        allow_empty_data=False,
        hash_data_files=False,
        price_data_root=str(Path(getattr(args, "price_data_root", None) or project_root).expanduser().resolve()),
        validate_price_data=True,
        scan_price_data=True,
        max_scan_rows=None,
        run_strategy=True,
        news_data_root=str(Path(getattr(args, "news_data_root", None) or project_root).expanduser().resolve()),
        disable_news_filter=bool(args.disable_news_filter),
        disable_news_forced_close=bool(args.disable_news_forced_close),
        news_source_timezone=None,
    )


def main() -> int:
    args = parse_args()
    try:
        engine_args = build_engine_args(args)
        summary = backtest_engine.run_engine(engine_args)
    except (
        backtest_engine.EngineError,
        backtest_engine.PriceDataError,
        backtest_engine.NewsDataError,
        backtest_engine.ExcelWriterError,
        backtest_engine.WorkbookConfigError,
    ) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    if engine_args.output_workbook:
        summary["output_workbook"] = engine_args.output_workbook
    print(json.dumps(summary, indent=2, sort_keys=True, default=to_jsonable))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
