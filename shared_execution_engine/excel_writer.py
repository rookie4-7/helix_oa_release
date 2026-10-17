from __future__ import annotations

import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet


RUN_LOG_HEADER_ROW = 3
RUN_LOG_FIRST_ROW = 4
RUN_LOG_LAST_ROW = 203
TRADE_LOG_HEADER_ROW = 3
TRADE_LOG_FIRST_ROW = 4
TRADE_LOG_LAST_ROW = 5000
DATA_CATALOG_HEADER_ROW = 3
DATA_CATALOG_FIRST_ROW = 4
DATA_CATALOG_LAST_ROW = 35
DATA_QUALITY_HEADER_ROW = 3
DATA_QUALITY_FIRST_ROW = 4
DATA_QUALITY_LAST_ROW = 35
SIGNAL_MODEL_HEADER_ROW = 4
SIGNAL_MODEL_FIRST_ROW = 5
SIGNAL_MODEL_LAST_ROW = 36
STRATEGY_PARAMS_HEADER_ROW = 3
STRATEGY_PARAMS_FIRST_ROW = 4
STRATEGY_PARAMS_LAST_ROW = 110
ANALYSIS_HEADER_ROW = 3
ANALYSIS_FIRST_ROW = 4
ANALYSIS_LAST_ROW = 35
YEAR_ANALYSIS_HEADER_ROW = 3
YEAR_ANALYSIS_FIRST_ROW = 4
YEAR_ANALYSIS_LAST_ROW = 2000
IS_REVIEW_HEADER_ROW = 3
IS_REVIEW_FIRST_ROW = 4
IS_REVIEW_LAST_ROW = 500
PARAMETER_SNAPSHOT_HEADER_ROW = 3
PARAMETER_SNAPSHOT_FIRST_ROW = 4
PARAMETER_SNAPSHOT_LAST_ROW = 500
IS_COMPARISON_HEADER_ROW = 3
IS_COMPARISON_FIRST_ROW = 4
IS_COMPARISON_LAST_ROW = 500
EQUITY_CURVE_HEADER_ROW = 3
EQUITY_CURVE_FIRST_ROW = 4
EQUITY_CURVE_LAST_ROW = 10000
SIGNAL_AUDIT_HEADER_ROW = 3
SIGNAL_AUDIT_FIRST_ROW = 4
SIGNAL_AUDIT_LAST_ROW = 20000
OOS_RESULT_REVIEW_HEADER_ROW = 3
OOS_RESULT_REVIEW_FIRST_ROW = 4
OOS_RESULT_REVIEW_LAST_ROW = 500

TITLE_FILL = "FF17324D"
SECTION_FILL = "FF008C95"
HEADER_FILL = "FFEAF3F8"
HEADER_FONT = "FF1F2933"
WHITE_FONT = "FFFFFFFF"
GRID_BORDER = "FFD9E2EA"

NUMBER_FORMATS = {
    "Run_DateTime": "yyyy-mm-dd hh:mm",
    "Locked_At": "yyyy-mm-dd hh:mm",
    "Win_Rate": "0.0%",
    "Year_Consistency": "0.0%",
    "Trades": "#,##0",
    "Wins": "#,##0",
    "Losses": "#,##0",
    "News_Blocked_Signals": "#,##0",
    "Valid_Signals": "#,##0",
    "Invalid_Signals": "#,##0",
    "No_Fill": "#,##0",
    "Blocked": "#,##0",
    "Forced_Closes": "#,##0",
    "Year": "0",
    "Net_R": "0.00",
    "Avg_R": "0.00",
    "Max_Drawdown_R": "0.00",
    "Trade_Based_Sharpe": "0.00",
    "Annualised_Trade_Sharpe": "0.00",
    "Daily_Sharpe": "0.00",
    "Annualised_Daily_Sharpe": "0.00",
    "Profit_Factor": "0.00",
    "Best_R": "0.00",
    "Worst_R": "0.00",
    "R_Multiple": "0.00",
    "Cumulative_R": "0.00",
    "Equity_Peak_R": "0.00",
    "Drawdown_R": "0.00",
    "Cumulative_Return_Pct": "0.0%",
    "PnL_Pct": "0.0%",
    "Entry_DateTime": "yyyy-mm-dd hh:mm",
    "Exit_DateTime": "yyyy-mm-dd hh:mm",
    "Signal_DateTime": "yyyy-mm-dd hh:mm",
    "Reviewed_At": "yyyy-mm-dd hh:mm",
    "Largest_Body_Pips": "0.0",
    "Entry": "0.00000",
    "Stop_Loss": "0.00000",
    "Take_Profit": "0.00000",
    "Run_Count": "#,##0",
    "Trade_Number": "#,##0",
    "Source_Row": "#,##0",
    "Run_Start": "#,##0",
    "Run_End": "#,##0",
    "OOS_Year": "0",
}

YEAR_ANALYSIS_HEADERS = [
    "Year_Record_ID",
    "Run_ID",
    "Window_ID",
    "Sample_Type",
    "Year",
    "Active_For_Analysis",
    "Official_OOS_Run",
    "Rule_Version",
    "Engine_Version",
    "Data_Version",
    "Trades",
    "Wins",
    "Losses",
    "Win_Rate",
    "Net_R",
    "Avg_R",
    "Trade_Based_Sharpe",
    "Annualised_Trade_Sharpe",
    "Daily_Sharpe",
    "Annualised_Daily_Sharpe",
    "Max_Drawdown_R",
    "Profit_Factor",
    "Best_R",
    "Worst_R",
    "News_Blocked_Signals",
    "Valid_Signals",
    "Invalid_Signals",
    "No_Fill",
    "Blocked",
    "Forced_Closes",
    "Notes",
]

IS_RUN_REVIEW_HEADERS = [
    "Review_Record_ID",
    "Window_ID",
    "Candidate_IS_Run_ID",
    "Run_DateTime",
    "IS_Period",
    "Keep_Reject",
    "Reason",
    "Selected_For_OOS",
    "Locked_At",
    "Locked_By",
    "Rule_Version",
    "Rule_Config_Hash",
    "Engine_Config_Hash",
    "Trades",
    "Net_R",
    "Max_Drawdown_R",
    "Trade_Based_Sharpe",
    "Annualised_Trade_Sharpe",
    "Daily_Sharpe",
    "Annualised_Daily_Sharpe",
    "Profit_Factor",
    "Year_Consistency",
    "Review_Status",
    "Notes",
]

PARAMETER_SNAPSHOT_HEADERS = [
    "Snapshot_ID",
    "Run_ID",
    "Window_ID",
    "Sample_Type",
    "Run_DateTime",
    "Rule_Version",
    "Rule_Set_ID",
    "Rule_Config_Hash",
    "Engine_Version",
    "Engine_Config_Hash",
    "Data_Version",
    "Key_Params",
    "Strategy_Params_Hash",
    "Strategy_Params_Path",
    "Rule_Config_Path",
    "Engine_Config_Path",
    "Archive_Path",
    "Notes",
]

IS_RUN_COMPARISON_HEADERS = [
    "Comparison_Record_ID",
    "Window_ID",
    "Run_ID",
    "Run_DateTime",
    "IS_Period",
    "Active_For_Analysis",
    "Keep_Reject",
    "Selected_For_OOS",
    "Rule_Version",
    "Rule_Config_Hash",
    "Engine_Config_Hash",
    "Data_Version",
    "Trades",
    "Wins",
    "Losses",
    "Win_Rate",
    "Net_R",
    "Max_Drawdown_R",
    "Trade_Based_Sharpe",
    "Annualised_Trade_Sharpe",
    "Daily_Sharpe",
    "Annualised_Daily_Sharpe",
    "Profit_Factor",
    "Best_R",
    "Worst_R",
    "News_Blocked_Signals",
    "Valid_Signals",
    "Invalid_Signals",
    "Year_Consistency",
    "Review_Notes",
]

EQUITY_CURVE_HEADERS = [
    "Equity_Record_ID",
    "Run_ID",
    "Window_ID",
    "Sample_Type",
    "Trade_Number",
    "Trade_ID",
    "Signal_ID",
    "Entry_DateTime",
    "Exit_DateTime",
    "Direction",
    "Setup_Tag",
    "Outcome",
    "Exit_Reason",
    "R_Multiple",
    "PnL_Pct",
    "Cumulative_R",
    "Equity_Peak_R",
    "Drawdown_R",
    "Cumulative_Return_Pct",
    "Notes",
]

SIGNAL_AUDIT_HEADERS = [
    "Audit_Record_ID",
    "Run_ID",
    "Window_ID",
    "Sample_Type",
    "Signal_ID",
    "Signal_DateTime",
    "Direction",
    "Setup_Type",
    "Status",
    "Reject_Category",
    "Reason",
    "Obstacle_Hit",
    "News_Blocked",
    "News_Event_IDs",
    "News_Events",
    "Run_Count",
    "Largest_Body_Pips",
    "Entry_Frozen",
    "Entry",
    "Stop_Loss",
    "Take_Profit",
    "Source_Row",
    "Run_Start",
    "Run_End",
    "Signal_Text",
    "Extension_Freeze",
    "Extension_Freeze_Reason",
    "Extension_Freeze_Touch_Time",
    "Extension_Freeze_Candidate_Time",
    "Extension_Freeze_Line_Version",
    "Context_Filter_Blocked",
    "Context_Filter_Reason",
    "Adverse_Asia_Reversal_Sell_Filter_Triggered",
    "Active_Price_Cluster_Stop_Filter_Triggered",
    "Adverse_Order_Block_Stop_Filter_Triggered",
    "Pre_Run_Structure_Direction",
    "Previous_Day_Signed_PFE_Trade_Direction",
    "Pre_Run_Asia_Signed_PFE_Trade_Direction",
    "Adverse_Asia_Reversal_Sell",
    "Active_Cluster_Present_Ahead",
    "Active_Cluster_Count_Ahead",
    "Active_Terminal_Cluster_ID",
    "Active_Terminal_Cluster_Low",
    "Active_Terminal_Cluster_High",
    "Active_Terminal_Cluster_Source_Move_ID",
    "Active_Terminal_Cluster_Source_Move_Start_Time",
    "Active_Terminal_Cluster_Source_Move_End_Time",
    "Active_Cluster_Interaction",
    "Run_Stopped_At_Active_Cluster",
    "Run_Cleared_Active_Cluster_Count",
    "Active_Order_Block_Count_Ahead",
    "Terminal_Order_Block_ID",
    "Terminal_Order_Block_Low",
    "Terminal_Order_Block_High",
    "Terminal_Order_Block_Source_Move_ID",
    "Terminal_Order_Block_Source_Move_Start_Time",
    "Terminal_Order_Block_Source_Move_End_Time",
    "Order_Block_Broad_Interaction",
    "Order_Block_Reached_Without_Clean_Close",
    "Run_Cleared_Order_Block_Count",
    "Adverse_Order_Block_Stop",
]

OOS_RESULT_REVIEW_HEADERS = [
    "Review_Record_ID",
    "Window_ID",
    "OOS_Run_ID",
    "Linked_IS_Run_ID",
    "Run_DateTime",
    "OOS_Period",
    "OOS_Year",
    "Official_OOS_Run",
    "OOS_Locked",
    "Review_Status",
    "Pass_Fail",
    "Continue_To_Next_Window",
    "Decision_Reason",
    "Reviewed_At",
    "Reviewed_By",
    "Rule_Version",
    "Rule_Config_Hash",
    "Engine_Config_Hash",
    "Trades",
    "Net_R",
    "Max_Drawdown_R",
    "Trade_Based_Sharpe",
    "Annualised_Trade_Sharpe",
    "Daily_Sharpe",
    "Annualised_Daily_Sharpe",
    "Profit_Factor",
    "Year_Consistency",
    "Notes",
]

STRATEGY_PARAMS_HEADERS = [
    "Category",
    "Parameter",
    "Value",
    "Type",
    "Unit",
    "Description",
]

WORKFLOW_COLUMN_WIDTHS = {
    "Category": 22,
    "Parameter": 34,
    "Value": 18,
    "Type": 12,
    "Unit": 14,
    "Description": 78,
    "Review_Record_ID": 22,
    "Comparison_Record_ID": 24,
    "Snapshot_ID": 22,
    "Candidate_IS_Run_ID": 20,
    "Run_ID": 14,
    "Window_ID": 14,
    "Run_DateTime": 22,
    "IS_Period": 24,
    "Rule_Config_Hash": 34,
    "Engine_Config_Hash": 34,
    "Strategy_Params_Hash": 34,
    "Data_Version": 42,
    "Key_Params": 72,
    "Strategy_Params_Path": 46,
    "Rule_Config_Path": 42,
    "Engine_Config_Path": 42,
    "Archive_Path": 42,
    "Reason": 30,
    "Review_Notes": 30,
    "Notes": 30,
    "Equity_Record_ID": 24,
    "Trade_ID": 18,
    "Signal_ID": 16,
    "Outcome": 14,
    "Exit_Reason": 18,
    "Audit_Record_ID": 24,
    "Signal_DateTime": 22,
    "Reject_Category": 18,
    "News_Event_IDs": 34,
    "News_Events": 52,
    "Signal_Text": 24,
    "OOS_Run_ID": 16,
    "Linked_IS_Run_ID": 18,
    "OOS_Period": 24,
    "Pass_Fail": 14,
    "Continue_To_Next_Window": 24,
    "Decision_Reason": 36,
    "Reviewed_By": 18,
}


def _strategy_param(
    category: str,
    parameter: str,
    value: Any,
    value_type: str,
    unit: str,
    description: str,
) -> dict[str, Any]:
    return {
        "Category": category,
        "Parameter": parameter,
        "Value": value,
        "Type": value_type,
        "Unit": unit,
        "Description": description,
    }


DEFAULT_STRATEGY_PARAM_ROWS: list[dict[str, Any]] = [
    _strategy_param("Setup switch", "use_asia_buy_setup", "Yes", "Bool", "Yes/No", "Allow Asia buy setups."),
    _strategy_param("Setup switch", "use_asia_sell_setup", "Yes", "Bool", "Yes/No", "Allow Asia sell setups."),
    _strategy_param("Setup switch", "use_daily_swing_buy_setup", "Yes", "Bool", "Yes/No", "Allow daily-swing buy setups."),
    _strategy_param("Setup switch", "use_daily_swing_sell_setup", "Yes", "Bool", "Yes/No", "Allow daily-swing sell setups."),
    _strategy_param("Setup switch", "use_consolidation_buy_setup", "Yes", "Bool", "Yes/No", "Allow consolidation buy setups."),
    _strategy_param("Setup switch", "use_consolidation_sell_setup", "Yes", "Bool", "Yes/No", "Allow consolidation sell setups."),
    _strategy_param("Setup switch", "use_internal_structure_buy_setup", "No", "Bool", "Yes/No", "Allow internal body-cluster BOS buy setups."),
    _strategy_param("Setup switch", "use_internal_structure_sell_setup", "No", "Bool", "Yes/No", "Allow internal body-cluster BOS sell setups."),
    _strategy_param("Internal BOS", "internal_structure_lookback_hours", 6.0, "Float", "hours", "Lookback window used to find internal body-high/body-low clusters."),
    _strategy_param("Internal BOS", "internal_structure_level_tolerance_pips", 3.0, "Float", "pips", "Maximum distance between body extremes in the internal level cluster."),
    _strategy_param("Internal BOS", "internal_structure_min_touches", 2, "Int", "touches", "Minimum body-extreme touches required to form an internal shelf."),
    _strategy_param("Internal BOS", "internal_structure_min_spacing_bars", 4, "Int", "15m bars", "Minimum spacing between the first and last internal shelf touches."),
    _strategy_param("Internal BOS", "internal_structure_max_shelf_body_range_pips", 60.0, "Float", "pips", "Maximum body range allowed in the internal shelf lookback window."),
    _strategy_param("Internal BOS", "use_internal_structure_asia_move_filter", "Yes", "Bool", "Yes/No", "Require Buy IBOS after an Asia body drop and Sell IBOS after an Asia body rise."),
    _strategy_param("Internal BOS", "internal_structure_min_asia_body_move_pct", 0.7, "Float", "%", "Minimum directional Asia body move required before IBOS is allowed."),
    _strategy_param("Setup switch", "use_buy_b_setup", "Yes", "Bool", "Yes/No", "Allow Buy B combination setups."),
    _strategy_param("Setup switch", "use_sell_b_setup", "Yes", "Bool", "Yes/No", "Allow Sell B combination setups."),
    _strategy_param("Filter switch", "use_buy_b_asia_range_filter", "Yes", "Bool", "Yes/No", "Require Buy B setups to stay within the configured Asia body range cap."),
    _strategy_param("Filter setting", "buy_b_max_asia_body_range_pips", 50.0, "Float", "pips", "Maximum Asia midpoint body range allowed for Buy B setups."),
    _strategy_param("Filter switch", "use_first_london_break_filter", "Yes", "Bool", "Yes/No", "Require the first eligible London break before setup signals."),
    _strategy_param("Filter switch", "use_imbalance_validation", "Yes", "Bool", "Yes/No", "Require the setup candle body to meet min_imbalance_pips."),
    _strategy_param("Filter switch", "use_pivot_obstacles", "Yes", "Bool", "Yes/No", "Block setups into active pivot obstacles."),
    _strategy_param("Filter switch", "use_consolidation_obstacles", "Yes", "Bool", "Yes/No", "Block setups into active consolidation obstacles."),
    _strategy_param("Filter switch", "use_obstacle_strength_override", "Yes", "Bool", "Yes/No", "Allow very strong imbalance runs to override selected obstacle blocks."),
    _strategy_param("Filter setting", "obstacle_override_min_run_range_ratio", 0.70, "Float", "ratio", "Minimum imbalance-run body range divided by Asia-to-pre-run body range required for obstacle override."),
    _strategy_param("Filter setting", "obstacle_override_min_largest_body_pips", 12.0, "Float", "pips", "Minimum largest imbalance candle body required for obstacle override."),
    _strategy_param("Filter switch", "obstacle_override_require_setup_confluence", "Yes", "Bool", "Yes/No", "Require the obstacle override run to have both Asia and daily-swing setup confirmation."),
    _strategy_param("Filter switch", "obstacle_override_apply_to_pivot", "Yes", "Bool", "Yes/No", "Allow the strength override to bypass pivot obstacle blocks."),
    _strategy_param("Filter switch", "obstacle_override_apply_to_consolidation", "No", "Bool", "Yes/No", "Allow the strength override to bypass consolidation obstacle blocks."),
    _strategy_param("Filter switch", "use_trendline_obstacles", "No", "Bool", "Yes/No", "Block setups whose run endpoint lands in an active trendline obstacle zone."),
    _strategy_param("Filter setting", "trendline_obstacle_pips", 1.0, "Float", "pips", "Half-width of the active trendline obstacle zone around the projected trendline price."),
    _strategy_param("Filter setting", "trendline_validation_tolerance_pips", 1.0, "Float", "pips", "Body-break tolerance used when validating and deactivating trendline obstacles."),
    _strategy_param("Context filter switch", "use_adverse_asia_reversal_sell_filter", "No", "Bool", "Yes/No", "Block Sell setups when protected structure and previous-day PFE are bullish while pre-run Asia PFE is bearish."),
    _strategy_param("Context filter switch", "use_active_price_cluster_stop_filter", "No", "Bool", "Yes/No", "Block a completed imbalance run that reaches an active directional strong-move price cluster without closing beyond it."),
    _strategy_param("Context filter switch", "use_adverse_order_block_stop_filter", "No", "Bool", "Yes/No", "Block an adverse completed imbalance run that reaches an active pivot-launched order block without closing beyond it."),
    _strategy_param("Context filter setting", "context_structure_confirm_bars", 5, "Int", "bars", "Left/right body-pivot confirmation bars for protected structure."),
    _strategy_param("Context filter setting", "context_structure_break_buffer_pips", 0.1, "Float", "pips", "Close-break buffer for protected body-pivot structure."),
    _strategy_param("Context filter setting", "context_strong_move_min_net_pips", 35.0, "Float", "pips", "Minimum net close displacement for a pivot-launched strong move."),
    _strategy_param("Context filter setting", "context_strong_move_min_pfe", 65.0, "Float", "PFE", "Minimum path efficiency for a pivot-launched strong move."),
    _strategy_param("Context filter setting", "context_strong_move_min_intervals", 4, "Int", "M15 intervals", "Minimum strong-move window length; four intervals equals one hour."),
    _strategy_param("Context filter setting", "context_strong_move_max_intervals", 16, "Int", "M15 intervals", "Maximum strong-move window length; sixteen intervals equals four hours."),
    _strategy_param("Context filter setting", "context_strong_move_local_pivot_lookback_bars", 4, "Int", "bars", "Bars used to confirm that the launch reversal pair is at a local price pivot."),
    _strategy_param("Context filter setting", "context_strong_move_episode_max_bars", 4, "Int", "bars", "Maximum launch spacing for collapsing overlapping strong-move windows into one impulse."),
    _strategy_param("Context filter setting", "context_structural_tolerance_pips", 3.0, "Float", "pips", "Price tolerance for collapsing overlapping cluster and order-block areas."),
    _strategy_param("Context filter setting", "context_order_block_gap_tolerance_pips", 3.0, "Float", "pips", "Maximum outside-wick gap allowed at the order-block launch pair."),
    _strategy_param("Context filter setting", "context_cluster_min_bars", 3, "Int", "bars", "Minimum candles in a strong-move price cluster."),
    _strategy_param("Context filter setting", "context_cluster_max_body_width_pips", 10.0, "Float", "pips", "Maximum total body-price width of a strong-move price cluster."),
    _strategy_param("Context filter setting", "context_cluster_max_net_progress_pips", 5.0, "Float", "pips", "Maximum net close progress inside a strong-move price cluster."),
    _strategy_param("Context filter setting", "context_cluster_max_pfe", 50.0, "Float", "PFE", "Maximum path efficiency inside a strong-move price cluster."),
    _strategy_param("Filter switch", "block_same_direction_after_first_signal_per_day", "No", "Bool", "Yes/No", "Block repeat same-direction setup signals after the first valid tradable signal per London day."),
    _strategy_param("Filter switch", "block_same_direction_after_first_trade_per_day", "No", "Bool", "Yes/No", "Block repeat same-direction trades after the first filled trade in that direction per London day."),
    _strategy_param("Filter switch", "only_first_filled_trade_per_day", "Yes", "Bool", "Yes/No", "Block every later filled trade after the first filled trade of each London day."),
    _strategy_param("Exit switch", "force_close_at_18_london", "Yes", "Bool", "Yes/No", "Force open trades closed at 18:00 Europe/London."),
    _strategy_param("News switch", "use_news_filter", "Yes", "Bool", "Yes/No", "Legacy switch name. Invalidate imbalance signals when the run candle window contains high-impact GBP/USD news."),
    _strategy_param("News switch", "use_news_forced_close", "Yes", "Bool", "Yes/No", "When the news layer is enabled, close open trades before same-day high-impact GBP/USD news."),
    _strategy_param("News", "news_forced_close_minutes_before", 15, "Int", "minutes", "Minutes before the news event to force-close any still-open same-day trade."),
    _strategy_param("Break setting", "structure_break_pips", 0.1, "Float", "pips", "Minimum structure break distance."),
    _strategy_param("Break setting", "setup_break_pips", 2.0, "Float", "pips", "Minimum setup break distance."),
    _strategy_param("Imbalance", "min_imbalance_pips", 5.0, "Float", "pips", "Minimum body size for a valid imbalance signal."),
    _strategy_param("Imbalance", "max_imbalance_candles", 5, "Int", "candles", "Maximum candles in a valid imbalance run."),
    _strategy_param("Entry", "entry_retracement", 0.618, "Float", "ratio", "Retracement level used for entry line calculation."),
    _strategy_param("Entry", "entry_line_hours", 1.0, "Float", "legacy", "Legacy field. Live entry lines now remain active until 12:00 Europe/London."),
    _strategy_param("Risk", "reverse_trade_direction", "No", "Bool", "Yes/No", "Reverse executed trade direction while keeping the original setup signal and fib entry line."),
    _strategy_param("Risk", "use_flexible_sl_tp", "No", "Bool", "Yes/No", "Use manual-style body/wick structure stop and R-multiple take profit instead of fixed pip distances."),
    _strategy_param("Risk", "flexible_stop_fib_level", 1.025, "Float", "legacy", "Legacy field. Flexible SL now uses body/wick structure distance plus buffer and risk clamps."),
    _strategy_param("Risk", "flexible_structure_buffer_pips", 2.9, "Float", "pips", "Flexible SL buffer added after averaging body and wick structure distance."),
    _strategy_param("Risk", "flexible_min_stop_loss_pips", 10.3, "Float", "pips", "Minimum flexible SL distance after body/wick structure calculation."),
    _strategy_param("Risk", "flexible_max_stop_loss_pips", 14.9, "Float", "pips", "Maximum flexible SL distance after body/wick structure calculation."),
    _strategy_param("Risk", "flexible_take_profit_r_multiple", 3.0, "Float", "R", "Flexible TP distance as a multiple of the entry-to-stop distance."),
    _strategy_param("Risk", "fixed_stop_loss_pips", 12.5, "Float", "pips", "Fixed stop loss distance."),
    _strategy_param("Risk", "fixed_take_profit_pips", 37.5, "Float", "pips", "Fixed take profit distance."),
    _strategy_param("Risk", "risk_per_trade", 0.01, "Float", "decimal", "Risk per trade used for equity curve percentage output."),
    _strategy_param("Exit", "max_exit_bars_after_entry", 384, "Int", "1m bars", "Safety cap for maximum 1-minute exit search after entry."),
    _strategy_param("Exit switch", "use_midpoint_tp_protective_exit", "No", "Bool", "Yes/No", "After midpoint touches TP, protect the trade with a fallback R exit if executable bid/ask does not confirm TP."),
    _strategy_param("Exit", "midpoint_tp_protective_exit_r_multiple", 2.0, "Float", "R", "Fallback R exit level armed after midpoint touches the normal TP."),
    _strategy_param("Execution", "execution_starts_after_signal_close", "Yes", "Bool", "Yes/No", "Start 1m execution only after the 15m signal candle has closed."),
    _strategy_param("Execution", "entry_fill_trigger_price_mode", "midpoint", "Text", "bid_ask/midpoint", "Use midpoint to trigger entry-line fills, then execute at bid/ask; use bid_ask for legacy executable-price line fills."),
    _strategy_param("Pivot", "left_bars", 5, "Int", "bars", "Left bars used for pivot detection."),
    _strategy_param("Pivot", "right_bars", 5, "Int", "bars", "Right bars used for pivot confirmation."),
    _strategy_param("Consolidation", "min_consolidation_hours", 2.0, "Float", "hours", "Minimum consolidation window length."),
    _strategy_param("Consolidation", "max_consolidation_body_range_pips", 10.0, "Float", "pips", "Maximum midpoint body range for consolidation."),
    _strategy_param("Consolidation", "similar_level_pips", 0.5, "Float", "pips", "Tolerance for similar highs/lows inside consolidation."),
    _strategy_param("Consolidation", "similar_level_min_spacing", 2, "Int", "bars", "Minimum bar spacing between similar consolidation levels."),
    _strategy_param("News", "news_blackout_minutes_before", 30, "Int", "legacy", "Legacy blackout field. Signal invalidation now checks whether news occurs inside the imbalance run candle window."),
    _strategy_param("News", "news_blackout_minutes_after", 30, "Int", "legacy", "Legacy blackout field. Signal invalidation now checks whether news occurs inside the imbalance run candle window."),
    _strategy_param("News", "news_currencies", "GBP,USD", "Text", "CSV", "News currencies included by the filter."),
    _strategy_param("News", "news_impacts", "High", "Text", "CSV", "News impact levels included by the filter."),
    _strategy_param("News", "news_source_timezone", "Europe/London", "Text", "timezone", "Timezone for naive news calendar times."),
]


def _signal_rule(
    rule_id: str,
    group: str,
    name: str,
    side: str,
    setup_type: str,
    combination_id: str,
    timeframe: str,
    input_source: str,
    condition: str,
    priority: int,
    enabled: str = "Yes",
    notes: str = "",
) -> dict[str, Any]:
    return {
        "Rule_ID": rule_id,
        "Rule_Group": group,
        "Rule_Name": name,
        "Side": side,
        "Setup_Type": setup_type,
        "Combination_ID": combination_id,
        "Timeframe": timeframe,
        "Input_Source": input_source,
        "Condition": condition,
        "Enabled": enabled,
        "Applies_To": "IS/OOS",
        "Priority": priority,
        "Version": "V1",
        "OOS_Lock": "No",
        "Notes": notes,
    }


DEFAULT_SIGNAL_MODEL_ROWS: list[dict[str, Any]] = [
    _signal_rule(
        "CANDLE_SESSION",
        "Foundation",
        "Candle body, pip, and London session model",
        "Both",
        "All",
        "FOUNDATION",
        "15m",
        "MBA price",
        "Use body prices for structure; evaluate Asia 00:00-07:44, pre-London structure context 07:45-07:59, London signals 08:00-11:59 Europe/London.",
        10,
        notes="From Shared_Candle_Session_Rules.docx.",
    ),
    _signal_rule(
        "ASIA_RANGE",
        "Setup State",
        "Asia body range tracking",
        "Both",
        "Asia",
        "ASIA_STATE",
        "15m",
        "MBA price",
        "Track Asia body high/low during 00:00-07:44; Asia breaks from 07:45-07:59 can arm context but cannot create entry lines.",
        20,
    ),
    _signal_rule(
        "ASIA_BUY",
        "Setup",
        "A Buy structure and setup break",
        "Buy",
        "A Buy",
        "A_BUY",
        "15m",
        "MBA price",
        "Bull candle breaks Asia high by structure_break, previous body was at/below Asia high, then same candle must exceed setup_break for setup.",
        30,
    ),
    _signal_rule(
        "ASIA_SELL",
        "Setup",
        "A Sell structure and setup break",
        "Sell",
        "A Sell",
        "A_SELL",
        "15m",
        "MBA price",
        "Bear candle breaks Asia low by structure_break, previous body was at/above Asia low, then same candle must exceed setup_break for setup.",
        40,
    ),
    _signal_rule(
        "DS_PIVOTS",
        "Setup State",
        "Daily Swing pivot level construction",
        "Both",
        "DS",
        "DS_STATE",
        "15m",
        "MBA price",
        "Use body pivots with left/right bars inside the 00:00-11:59 DS window; highest DS high and lowest DS low reset each London day.",
        50,
    ),
    _signal_rule(
        "DS_BUY",
        "Setup",
        "DS Buy structure and setup break",
        "Buy",
        "DS Buy",
        "DS_BUY",
        "15m",
        "MBA price",
        "Bull candle breaks active DS high by structure_break, previous body was at/below DS high, then same candle must exceed setup_break for setup.",
        60,
    ),
    _signal_rule(
        "DS_SELL",
        "Setup",
        "DS Sell structure and setup break",
        "Sell",
        "DS Sell",
        "DS_SELL",
        "15m",
        "MBA price",
        "Bear candle breaks active DS low by structure_break, previous body was at/above DS low, then same candle must exceed setup_break for setup.",
        70,
    ),
    _signal_rule(
        "CONSOL_READY",
        "Setup State",
        "Consolidation readiness",
        "Both",
        "Consolidation",
        "CONSOL_STATE",
        "15m",
        "MBA price",
        "Build prior-body box excluding current candle; require max range plus spaced similar highs/lows; compression and efficiency filters are optional.",
        80,
    ),
    _signal_rule(
        "CONSOL_BUY",
        "Setup",
        "Consolidation Buy setup",
        "Buy",
        "Consol.BUY",
        "CONSOL_BUY",
        "15m",
        "MBA price",
        "Requires similar highs, bull candle, raw break above consolidation high, and close above consolidation high by setup_break.",
        90,
    ),
    _signal_rule(
        "CONSOL_SELL",
        "Setup",
        "Consolidation Sell setup",
        "Sell",
        "Consol.SELL",
        "CONSOL_SELL",
        "15m",
        "MBA price",
        "Requires similar lows, bear candle, raw break below consolidation low, and close below consolidation low by setup_break.",
        100,
    ),
    _signal_rule(
        "FIRST_LONDON_BREAK_BUY",
        "Gate",
        "First London buy-side break lock",
        "Buy",
        "All",
        "FIRST_BREAK",
        "15m",
        "MBA price",
        "During 08:00-11:59, first buy-side structure break from Asia, DS, or Consolidation locks buy-side signal permission; 07:45-07:59 Asia/DS context can be inherited by the first valid same-side London imbalance.",
        110,
    ),
    _signal_rule(
        "FIRST_LONDON_BREAK_SELL",
        "Gate",
        "First London sell-side break lock",
        "Sell",
        "All",
        "FIRST_BREAK",
        "15m",
        "MBA price",
        "During 08:00-11:59, first sell-side structure break from Asia, DS, or Consolidation locks sell-side signal permission; 07:45-07:59 Asia/DS context can be inherited by the first valid same-side London imbalance.",
        120,
    ),
    _signal_rule(
        "IMBALANCE_RUN",
        "Validation",
        "Same-direction imbalance run construction",
        "Both",
        "All",
        "RUN",
        "15m",
        "MBA price",
        "Build the full consecutive body-direction run, not only candles after setup; doji resets the run.",
        130,
    ),
    _signal_rule(
        "RUN_VALIDATION",
        "Validation",
        "Base run validation",
        "Both",
        "All",
        "RUN",
        "15m",
        "MBA price",
        "Final signal requires run_count <= max_imbalance_candles and at least one setup candle body >= min_imbalance_pips.",
        140,
    ),
    _signal_rule(
        "BUY_B",
        "Combination",
        "Buy B state-machine confirmation",
        "Buy",
        "Buy B",
        "B_SIGNAL",
        "15m",
        "MBA price",
        "Buy B requires valid Asia and valid DS buy confirmation in the same run, or current Asia/DS plus the other prior confirmed in-session.",
        150,
        notes="Consolidation is not a B ingredient.",
    ),
    _signal_rule(
        "SELL_B",
        "Combination",
        "Sell B state-machine confirmation",
        "Sell",
        "Sell B",
        "B_SIGNAL",
        "15m",
        "MBA price",
        "Sell B requires valid Asia and valid DS sell confirmation in the same run, or current Asia/DS plus the other prior confirmed in-session.",
        160,
        notes="Consolidation is not a B ingredient.",
    ),
    _signal_rule(
        "FINAL_BUY_LABEL",
        "Validation",
        "Final buy-side signal labelling",
        "Buy",
        "All Buy",
        "FINAL_SIGNAL",
        "15m",
        "MBA price",
        "After run finish: output Buy B, otherwise A Buy and/or DS Buy, and append Consol.BUY if valid consolidation body exists.",
        170,
    ),
    _signal_rule(
        "FINAL_SELL_LABEL",
        "Validation",
        "Final sell-side signal labelling",
        "Sell",
        "All Sell",
        "FINAL_SIGNAL",
        "15m",
        "MBA price",
        "After run finish: output Sell B, otherwise A Sell and/or DS Sell, and append Consol.SELL if valid consolidation body exists.",
        180,
    ),
    _signal_rule(
        "PIVOT_OBSTACLES",
        "Filter",
        "Pivot support/resistance obstacles",
        "Both",
        "All",
        "OBSTACLE",
        "15m",
        "MBA price",
        "Confirmed body pivots create support/resistance zones; buy uses run_high against resistance and sell uses run_low against support.",
        190,
    ),
    _signal_rule(
        "CONSOL_OBSTACLES",
        "Filter",
        "Broken consolidation obstacles",
        "Both",
        "All",
        "OBSTACLE",
        "15m",
        "MBA price",
        "Broken consolidation boxes become future obstacle zones only for runs whose start is after the obstacle-created bar.",
        200,
    ),
    _signal_rule(
        "TRENDLINE_OBSTACLES",
        "Filter",
        "Trendline obstacles",
        "Both",
        "All",
        "OBSTACLE",
        "15m",
        "MBA price",
        "Optional off-by-default trendline zones from recent pivot points; block run endpoint inside +/- trendline_obstacle_pips after trendline validation.",
        210,
        enabled="No",
        notes="Default Pine input use_trendline_obstacles=false; validation uses trendline_validation_tolerance_pips.",
    ),
    _signal_rule(
        "LIVE_SIGNAL_GATE",
        "Live State",
        "Live signal and invalidation state",
        "Both",
        "All",
        "LIVE",
        "15m",
        "MBA price",
        "Live signals can appear before final run finish and are suppressed by live obstacle hits; final invalid labels require the signal not already frozen.",
        220,
    ),
    _signal_rule(
        "ENTRY_LINE",
        "Entry",
        "61.8% entry line calculation",
        "Both",
        "All",
        "ENTRY",
        "15m",
        "MBA price",
        "Calculate buy/sell entry from closed-candle run body high/low using entry_retracement; live line uses closed-candle basis.",
        230,
    ),
    _signal_rule(
        "ENTRY_FREEZE",
        "Entry",
        "Entry and signal freeze",
        "Both",
        "All",
        "ENTRY",
        "15m",
        "MBA price",
        "If wick high/low touches the live entry before run finish, freeze entry price/time/direction and do not rewrite the signal.",
        240,
    ),
    _signal_rule(
        "EXECUTION_SL_TP",
        "Execution",
        "Fixed SL/TP execution overlay",
        "Both",
        "All",
        "EXECUTION",
        "1m",
        "MBA price",
        "Backtest-only execution uses fixed 12.5 pip SL and 37.5 pip TP on lower timeframe bid/ask fills.",
        250,
    ),
    _signal_rule(
        "FORCE_CLOSE_18",
        "Execution",
        "18:00 London forced close",
        "Both",
        "All",
        "EXECUTION",
        "1m",
        "MBA price",
        "Close any still-open trade at or after 18:00 Europe/London on the same London entry date.",
        260,
    ),
    _signal_rule(
        "SAME_DIRECTION_BLOCK",
        "Execution",
        "Same-direction second filled trade block",
        "Both",
        "All",
        "EXECUTION",
        "1m",
        "MBA price",
        "After first filled trade for a London date, block later filled trades in the same direction; opposite direction remains eligible.",
        270,
    ),
    _signal_rule(
        "ONLY_FIRST_FILLED_TRADE_DAY",
        "Execution",
        "Only first filled trade per day",
        "Both",
        "All",
        "EXECUTION",
        "1m",
        "MBA price",
        "When enabled, after the first filled trade for a London date, block every later fill attempt regardless of direction.",
        275,
    ),
    _signal_rule(
        "NEWS_INVALIDATION",
        "Invalidation",
        "News-caused imbalance invalidation",
        "Both",
        "All",
        "NEWS",
        "Calendar",
        "News calendar",
        "Invalidate an imbalance signal when a high-impact GBP/USD news event occurs inside the imbalance run candle window.",
        280,
        notes="Project extension beyond the original Pine docs.",
    ),
    _signal_rule(
        "NEWS_FORCED_CLOSE",
        "Execution",
        "Pre-news forced close",
        "Both",
        "All",
        "NEWS",
        "1m",
        "News calendar + MBA price",
        "If a trade is still open on the same London day as a high-impact GBP/USD event, close it 15 minutes before the release unless SL/TP hits first.",
        290,
        notes="Project extension beyond the original Pine docs.",
    ),
]


class ExcelWriterError(Exception):
    """Raised when the workbook cannot accept generated engine output."""


def table_headers(sheet: Worksheet, header_row: int) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col_idx in range(1, sheet.max_column + 1):
        value = sheet.cell(row=header_row, column=col_idx).value
        if isinstance(value, str) and value.strip():
            headers[value.strip()] = col_idx
    return headers


def next_run_ids(run_log: Worksheet, count: int) -> list[str]:
    pattern = re.compile(r"^RUN_(\d+)$")
    max_id = 0
    run_id_col = table_headers(run_log, RUN_LOG_HEADER_ROW).get("Run_ID", 1)
    for row_idx in range(RUN_LOG_FIRST_ROW, RUN_LOG_LAST_ROW + 1):
        value = run_log.cell(row=row_idx, column=run_id_col).value
        if not isinstance(value, str):
            continue
        match = pattern.match(value.strip())
        if match:
            max_id = max(max_id, int(match.group(1)))
    return [f"RUN_{idx:04d}" for idx in range(max_id + 1, max_id + count + 1)]


def empty_run_log_rows(run_log: Worksheet, count: int) -> list[int]:
    return empty_rows(
        sheet=run_log,
        id_column="Run_ID",
        header_row=RUN_LOG_HEADER_ROW,
        first_row=RUN_LOG_FIRST_ROW,
        last_row=RUN_LOG_LAST_ROW,
        count=count,
        label="Run_Log",
    )


def empty_trade_log_rows(trade_log: Worksheet, count: int) -> list[int]:
    return empty_rows(
        sheet=trade_log,
        id_column="Trade_ID",
        header_row=TRADE_LOG_HEADER_ROW,
        first_row=TRADE_LOG_FIRST_ROW,
        last_row=TRADE_LOG_LAST_ROW,
        count=count,
        label="Trade_Log",
    )


def ensure_default_signal_model(sheet: Worksheet) -> int:
    """Seed the rule inventory once, while leaving user-edited sheets alone."""
    if _sheet_has_data(sheet, SIGNAL_MODEL_HEADER_ROW, SIGNAL_MODEL_FIRST_ROW, SIGNAL_MODEL_LAST_ROW):
        return 0
    if len(DEFAULT_SIGNAL_MODEL_ROWS) > SIGNAL_MODEL_LAST_ROW - SIGNAL_MODEL_FIRST_ROW + 1:
        raise ExcelWriterError("Default Signal_Model inventory does not fit in the workbook table.")

    for offset, values in enumerate(DEFAULT_SIGNAL_MODEL_ROWS):
        write_named_row(
            sheet,
            SIGNAL_MODEL_HEADER_ROW,
            SIGNAL_MODEL_FIRST_ROW + offset,
            values,
            required=("Rule_ID",),
        )
    return len(DEFAULT_SIGNAL_MODEL_ROWS)


def ensure_analysis_metric_columns(workbook: Any) -> None:
    ensure_strategy_params_sheet(workbook)
    for sheet_name in ("IS_Analysis", "OOS_Analysis"):
        ensure_sample_analysis_columns(workbook[sheet_name])
    ensure_walk_forward_analysis_columns(workbook["WF_Window_Analysis"])
    ensure_dashboard_sharpe_rows(workbook["Dashboard"])
    ensure_year_analysis_sheet(workbook)
    ensure_is_run_review_sheet(workbook)
    ensure_is_run_comparison_sheet(workbook)
    ensure_parameter_snapshot_sheet(workbook)
    ensure_equity_curve_sheet(workbook)
    ensure_signal_audit_sheet(workbook)
    ensure_oos_result_review_sheet(workbook)


def ensure_strategy_params_sheet(workbook: Any) -> Worksheet:
    sheet = ensure_table_sheet(
        workbook=workbook,
        sheet_name="Strategy_Params",
        title="Strategy Params",
        headers=STRATEGY_PARAMS_HEADERS,
        after_sheet="Config",
        header_row=STRATEGY_PARAMS_HEADER_ROW,
    )
    sheet["A2"] = "Edit the Value column only. Python reads this sheet before every strategy run and archives the exact snapshot."
    sheet["A2"].font = Font(color=HEADER_FONT, italic=True, size=10)
    sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 30

    headers = table_headers(sheet, STRATEGY_PARAMS_HEADER_ROW)
    parameter_col = headers.get("Parameter")
    value_col = headers.get("Value")
    type_col = headers.get("Type")
    if not parameter_col or not value_col or not type_col:
        raise ExcelWriterError("Strategy_Params is missing Parameter, Value, or Type headers.")

    existing_rows: dict[str, int] = {}
    for row_idx in range(STRATEGY_PARAMS_FIRST_ROW, max(sheet.max_row, STRATEGY_PARAMS_LAST_ROW) + 1):
        parameter = sheet.cell(row=row_idx, column=parameter_col).value
        if isinstance(parameter, str) and parameter.strip():
            existing_rows[parameter.strip()] = row_idx

    next_row = max(existing_rows.values(), default=STRATEGY_PARAMS_FIRST_ROW - 1) + 1
    for default_row in DEFAULT_STRATEGY_PARAM_ROWS:
        parameter = str(default_row["Parameter"])
        row_idx = existing_rows.get(parameter)
        if row_idx is None:
            if next_row > STRATEGY_PARAMS_LAST_ROW:
                raise ExcelWriterError("Default Strategy_Params inventory does not fit in the workbook table.")
            write_named_row(sheet, STRATEGY_PARAMS_HEADER_ROW, next_row, default_row, required=("Parameter", "Value"))
            apply_table_body_row_theme(sheet, STRATEGY_PARAMS_HEADERS, next_row)
            row_idx = next_row
            next_row += 1
        else:
            for header in ("Category", "Type", "Unit", "Description"):
                col_idx = headers.get(header)
                if col_idx is not None and sheet.cell(row=row_idx, column=col_idx).value in (None, ""):
                    sheet.cell(row=row_idx, column=col_idx, value=default_row[header])
            apply_table_body_row_theme(sheet, STRATEGY_PARAMS_HEADERS, row_idx)

    apply_strategy_params_validation(sheet)
    return sheet


def apply_strategy_params_validation(sheet: Worksheet) -> None:
    headers = table_headers(sheet, STRATEGY_PARAMS_HEADER_ROW)
    value_col = headers.get("Value")
    type_col = headers.get("Type")
    if not value_col or not type_col:
        return

    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    yes_no.error = "Enter Yes or No."
    yes_no.errorTitle = "Invalid switch value"
    yes_no.prompt = "Choose Yes or No."
    yes_no.promptTitle = "Strategy switch"
    sheet.add_data_validation(yes_no)
    for row_idx in range(STRATEGY_PARAMS_FIRST_ROW, STRATEGY_PARAMS_LAST_ROW + 1):
        if str(sheet.cell(row=row_idx, column=type_col).value or "").strip().lower() == "bool":
            yes_no.add(sheet.cell(row=row_idx, column=value_col))

    for row_idx in range(STRATEGY_PARAMS_FIRST_ROW, min(sheet.max_row, STRATEGY_PARAMS_LAST_ROW) + 1):
        value_cell = sheet.cell(row=row_idx, column=value_col)
        type_value = str(sheet.cell(row=row_idx, column=type_col).value or "").strip().lower()
        if type_value == "float":
            value_cell.number_format = "0.000"
        elif type_value == "int":
            value_cell.number_format = "0"
        elif type_value == "bool":
            value_cell.number_format = "@"


def ensure_year_analysis_sheet(workbook: Any) -> Worksheet:
    return ensure_table_sheet(
        workbook=workbook,
        sheet_name="Year_Analysis",
        title="Year Analysis",
        headers=YEAR_ANALYSIS_HEADERS,
        after_sheet="OOS_Analysis",
        header_row=YEAR_ANALYSIS_HEADER_ROW,
    )


def ensure_is_run_review_sheet(workbook: Any) -> Worksheet:
    return ensure_table_sheet(
        workbook=workbook,
        sheet_name="IS_Run_Review",
        title="IS Run Review",
        headers=IS_RUN_REVIEW_HEADERS,
        after_sheet="Year_Analysis",
        header_row=IS_REVIEW_HEADER_ROW,
    )


def ensure_is_run_comparison_sheet(workbook: Any) -> Worksheet:
    return ensure_table_sheet(
        workbook=workbook,
        sheet_name="IS_Run_Comparison",
        title="IS Run Comparison",
        headers=IS_RUN_COMPARISON_HEADERS,
        after_sheet="IS_Run_Review",
        header_row=IS_COMPARISON_HEADER_ROW,
    )


def ensure_parameter_snapshot_sheet(workbook: Any) -> Worksheet:
    return ensure_table_sheet(
        workbook=workbook,
        sheet_name="Parameter_Snapshot",
        title="Parameter Snapshot",
        headers=PARAMETER_SNAPSHOT_HEADERS,
        after_sheet="IS_Run_Comparison",
        header_row=PARAMETER_SNAPSHOT_HEADER_ROW,
    )


def ensure_equity_curve_sheet(workbook: Any) -> Worksheet:
    return ensure_table_sheet(
        workbook=workbook,
        sheet_name="Equity_Curve",
        title="Equity Curve",
        headers=EQUITY_CURVE_HEADERS,
        after_sheet="Parameter_Snapshot",
        header_row=EQUITY_CURVE_HEADER_ROW,
    )


def ensure_signal_audit_sheet(workbook: Any) -> Worksheet:
    return ensure_table_sheet(
        workbook=workbook,
        sheet_name="Signal_Audit",
        title="Signal Audit",
        headers=SIGNAL_AUDIT_HEADERS,
        after_sheet="Equity_Curve",
        header_row=SIGNAL_AUDIT_HEADER_ROW,
    )


def ensure_oos_result_review_sheet(workbook: Any) -> Worksheet:
    return ensure_table_sheet(
        workbook=workbook,
        sheet_name="OOS_Result_Review",
        title="OOS Result Review",
        headers=OOS_RESULT_REVIEW_HEADERS,
        after_sheet="Signal_Audit",
        header_row=OOS_RESULT_REVIEW_HEADER_ROW,
    )


def ensure_table_sheet(
    workbook: Any,
    sheet_name: str,
    title: str,
    headers: list[str],
    after_sheet: str,
    header_row: int,
) -> Worksheet:
    if sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
    else:
        index = workbook.sheetnames.index(after_sheet) + 1 if after_sheet in workbook.sheetnames else len(workbook.sheetnames)
        sheet = workbook.create_sheet(sheet_name, index)
        sheet["A1"] = title
        sheet.freeze_panes = f"A{header_row + 1}"

    sheet["A1"] = sheet["A1"].value or title
    sheet.freeze_panes = sheet.freeze_panes or f"A{header_row + 1}"
    for col_idx, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=col_idx, value=header)
        width = WORKFLOW_COLUMN_WIDTHS.get(
            header,
            14 if len(header) < 14 else min(len(header) + 2, 34),
        )
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
    apply_table_theme(sheet, headers=headers, header_row=header_row)
    return sheet


def apply_table_theme(sheet: Worksheet, headers: list[str], header_row: int) -> None:
    max_col = len(headers)
    title_fill = PatternFill("solid", fgColor=TITLE_FILL)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    title_font = Font(color=WHITE_FONT, bold=True, size=10)
    header_font = Font(color=HEADER_FONT, bold=True, size=10)
    body_font = Font(color=HEADER_FONT, size=10)
    border = Border(bottom=Side(style="thin", color=GRID_BORDER))

    sheet.row_dimensions[1].height = 22
    sheet.row_dimensions[header_row].height = 20
    for col_idx in range(1, max_col + 1):
        title_cell = sheet.cell(row=1, column=col_idx)
        title_cell.fill = title_fill
        title_cell.font = title_font
        title_cell.alignment = Alignment(vertical="center")

        header_cell = sheet.cell(row=header_row, column=col_idx)
        header_cell.fill = header_fill
        header_cell.font = header_font
        header_cell.border = border
        header_cell.alignment = Alignment(vertical="center", wrap_text=True)

    sheet["A1"] = sheet["A1"].value
    for row_idx in range(header_row + 1, min(sheet.max_row, header_row + 12) + 1):
        for col_idx in range(1, max_col + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top")
            cell.border = border

    _apply_number_formats(sheet, headers, header_row)


def _apply_number_formats(sheet: Worksheet, headers: list[str], header_row: int) -> None:
    last_row = max(sheet.max_row, header_row + 1)
    for col_idx, header in enumerate(headers, start=1):
        number_format = NUMBER_FORMATS.get(header)
        if not number_format:
            continue
        for row_idx in range(header_row + 1, last_row + 1):
            sheet.cell(row=row_idx, column=col_idx).number_format = number_format


def apply_table_body_row_theme(sheet: Worksheet, headers: list[str], row_idx: int) -> None:
    body_font = Font(color=HEADER_FONT, size=10)
    border = Border(bottom=Side(style="thin", color=GRID_BORDER))
    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row_idx, column=col_idx)
        cell.font = body_font
        cell.alignment = Alignment(vertical="top")
        cell.border = border
        if header in NUMBER_FORMATS:
            cell.number_format = NUMBER_FORMATS[header]


def ensure_sample_analysis_columns(sheet: Worksheet) -> None:
    headers = table_headers(sheet, ANALYSIS_HEADER_ROW)
    if "Annualised_Sharpe" in headers and "Annualised_Trade_Sharpe" not in headers:
        sheet.cell(
            row=ANALYSIS_HEADER_ROW,
            column=headers["Annualised_Sharpe"],
            value="Annualised_Trade_Sharpe",
        )

    headers = table_headers(sheet, ANALYSIS_HEADER_ROW)
    insert_at = headers.get("Max_Drawdown_R") or headers.get("Profit_Factor") or sheet.max_column + 1
    for name in ("Daily_Sharpe", "Annualised_Daily_Sharpe"):
        headers = table_headers(sheet, ANALYSIS_HEADER_ROW)
        if name in headers:
            continue
        sheet.insert_cols(insert_at)
        sheet.cell(row=ANALYSIS_HEADER_ROW, column=insert_at, value=name)
        _copy_column_style(sheet, max(1, insert_at - 1), insert_at)
        insert_at += 1


def ensure_trade_log_signal_columns(sheet: Worksheet) -> None:
    headers = table_headers(sheet, TRADE_LOG_HEADER_ROW)
    if "Signal_ID" in headers and "Signal_DateTime" in headers:
        return

    insert_at = headers.get("Entry_DateTime") or headers.get("Setup_Tag", sheet.max_column) + 1
    for name in ("Signal_ID", "Signal_DateTime"):
        headers = table_headers(sheet, TRADE_LOG_HEADER_ROW)
        if name in headers:
            continue
        sheet.insert_cols(insert_at)
        _copy_column_style(sheet, max(1, insert_at - 1), insert_at)
        sheet.cell(row=TRADE_LOG_HEADER_ROW, column=insert_at, value=name)
        if name in NUMBER_FORMATS:
            for row_idx in range(TRADE_LOG_HEADER_ROW + 1, max(sheet.max_row, TRADE_LOG_HEADER_ROW + 1) + 1):
                sheet.cell(row=row_idx, column=insert_at).number_format = NUMBER_FORMATS[name]
        insert_at += 1


def ensure_walk_forward_analysis_columns(sheet: Worksheet) -> None:
    headers = [
        "Window_ID",
        "IS_Period",
        "OOS_Year",
        "Selected_IS_Run_ID",
        "Official_OOS_Run_ID",
        "IS_Trades",
        "OOS_Trades",
        "IS_Net_R",
        "OOS_Net_R",
        "IS_Win_Rate",
        "OOS_Win_Rate",
        "IS_Trade_Sharpe",
        "OOS_Trade_Sharpe",
        "IS_Annualised_Trade_Sharpe",
        "OOS_Annualised_Trade_Sharpe",
        "IS_Daily_Sharpe",
        "OOS_Daily_Sharpe",
        "IS_Annualised_Daily_Sharpe",
        "OOS_Annualised_Daily_Sharpe",
        "Net_R_Decay",
        "Trade_Sharpe_Decay",
        "Daily_Sharpe_Decay",
        "Review_Status",
        "Notes",
    ]
    for col_idx, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=col_idx, value=header)

    for row_idx in range(4, 12):
        sheet.cell(row=row_idx, column=4, value=f'=IF(\'Walk_Forward\'!I{row_idx}="","",\'Walk_Forward\'!I{row_idx})')
        sheet.cell(row=row_idx, column=5, value=f'=IF(\'Walk_Forward\'!K{row_idx}="","",\'Walk_Forward\'!K{row_idx})')
        _write_wf_metric_formula(sheet, row_idx, 6, "IS_Analysis", "Trades", "D")
        _write_wf_metric_formula(sheet, row_idx, 7, "OOS_Analysis", "Trades", "E")
        _write_wf_metric_formula(sheet, row_idx, 8, "IS_Analysis", "Net_R", "D")
        _write_wf_metric_formula(sheet, row_idx, 9, "OOS_Analysis", "Net_R", "E")
        _write_wf_metric_formula(sheet, row_idx, 10, "IS_Analysis", "Win_Rate", "D")
        _write_wf_metric_formula(sheet, row_idx, 11, "OOS_Analysis", "Win_Rate", "E")
        _write_wf_metric_formula(sheet, row_idx, 12, "IS_Analysis", "Trade_Based_Sharpe", "D")
        _write_wf_metric_formula(sheet, row_idx, 13, "OOS_Analysis", "Trade_Based_Sharpe", "E")
        _write_wf_metric_formula(sheet, row_idx, 14, "IS_Analysis", "Annualised_Trade_Sharpe", "D")
        _write_wf_metric_formula(sheet, row_idx, 15, "OOS_Analysis", "Annualised_Trade_Sharpe", "E")
        _write_wf_metric_formula(sheet, row_idx, 16, "IS_Analysis", "Daily_Sharpe", "D")
        _write_wf_metric_formula(sheet, row_idx, 17, "OOS_Analysis", "Daily_Sharpe", "E")
        _write_wf_metric_formula(sheet, row_idx, 18, "IS_Analysis", "Annualised_Daily_Sharpe", "D")
        _write_wf_metric_formula(sheet, row_idx, 19, "OOS_Analysis", "Annualised_Daily_Sharpe", "E")
        sheet.cell(row=row_idx, column=20, value=f"=IF(H{row_idx}=0,0,I{row_idx}/H{row_idx}-1)")
        sheet.cell(row=row_idx, column=21, value=f"=IF(L{row_idx}=0,0,M{row_idx}/L{row_idx}-1)")
        sheet.cell(row=row_idx, column=22, value=f"=IF(P{row_idx}=0,0,Q{row_idx}/P{row_idx}-1)")
        if sheet.cell(row=row_idx, column=23).value in (None, ""):
            sheet.cell(row=row_idx, column=23, value="Pending")


def ensure_dashboard_sharpe_rows(sheet: Worksheet) -> None:
    values = {
        "A8": "OOS Trade Sharpe",
        "B8": '=IFERROR(AVERAGEIFS(\'OOS_Analysis\'!$P$4:$P$35,\'OOS_Analysis\'!$D$4:$D$35,"Yes",\'OOS_Analysis\'!$P$4:$P$35,"<>0"),0)',
        "A9": "OOS Annualised Trade Sharpe",
        "B9": '=IFERROR(AVERAGEIFS(\'OOS_Analysis\'!$Q$4:$Q$35,\'OOS_Analysis\'!$D$4:$D$35,"Yes",\'OOS_Analysis\'!$Q$4:$Q$35,"<>0"),0)',
        "A10": "OOS Daily Sharpe",
        "B10": '=IFERROR(AVERAGEIFS(\'OOS_Analysis\'!$R$4:$R$35,\'OOS_Analysis\'!$D$4:$D$35,"Yes",\'OOS_Analysis\'!$R$4:$R$35,"<>0"),0)',
        "A11": "OOS Annualised Daily Sharpe",
        "B11": '=IFERROR(AVERAGEIFS(\'OOS_Analysis\'!$S$4:$S$35,\'OOS_Analysis\'!$D$4:$D$35,"Yes",\'OOS_Analysis\'!$S$4:$S$35,"<>0"),0)',
        "F7": "IS Trade Sharpe",
        "G7": '=IFERROR(AVERAGEIFS(\'IS_Analysis\'!$P$4:$P$35,\'IS_Analysis\'!$D$4:$D$35,"Yes",\'IS_Analysis\'!$P$4:$P$35,"<>0"),0)',
        "F8": "IS Annualised Trade Sharpe",
        "G8": '=IFERROR(AVERAGEIFS(\'IS_Analysis\'!$Q$4:$Q$35,\'IS_Analysis\'!$D$4:$D$35,"Yes",\'IS_Analysis\'!$Q$4:$Q$35,"<>0"),0)',
        "F9": "IS Daily Sharpe",
        "G9": '=IFERROR(AVERAGEIFS(\'IS_Analysis\'!$R$4:$R$35,\'IS_Analysis\'!$D$4:$D$35,"Yes",\'IS_Analysis\'!$R$4:$R$35,"<>0"),0)',
        "F10": "IS Annualised Daily Sharpe",
        "G10": '=IFERROR(AVERAGEIFS(\'IS_Analysis\'!$S$4:$S$35,\'IS_Analysis\'!$D$4:$D$35,"Yes",\'IS_Analysis\'!$S$4:$S$35,"<>0"),0)',
        "F11": "Avg IS/OOS Net R Decay",
        "G11": "=IFERROR(AVERAGE('WF_Window_Analysis'!T4:T11),0)",
        "F12": "Avg Trade Sharpe Decay",
        "G12": "=IFERROR(AVERAGE('WF_Window_Analysis'!U4:U11),0)",
        "F13": "Avg Daily Sharpe Decay",
        "G13": "=IFERROR(AVERAGE('WF_Window_Analysis'!V4:V11),0)",
    }
    for coordinate, value in values.items():
        sheet[coordinate] = value


def write_data_catalog(sheet: Worksheet, manifest: dict[str, Any], now: datetime) -> None:
    files = list(manifest.get("files") or [])
    _require_capacity("Data_Catalog", len(files), DATA_CATALOG_FIRST_ROW, DATA_CATALOG_LAST_ROW)
    _clear_rows(sheet, DATA_CATALOG_FIRST_ROW, DATA_CATALOG_LAST_ROW)
    _report_choices(sheet, "M4:M35", "Pending,Validated,Review,Failed,Headers Checked,Partially Checked,Archived")

    for offset, item in enumerate(files):
        data_id = item.get("data_id")
        values = {
            "Data_ID": data_id,
            "Symbol": item.get("symbol"),
            "Timeframe": item.get("timeframe"),
            "Data_Type": item.get("data_type") or "MBA",
            "File_Path": item.get("file_path"),
            "File_Format": item.get("file_format"),
            "Start_Date": item.get("start_date"),
            "End_Date": item.get("end_date"),
            "Rows": item.get("rows"),
            "Timezone": item.get("timezone"),
            "Data_Version": item.get("data_version"),
            "Checksum": item.get("sha256") or item.get("checksum"),
            "Status": item.get("report_status", item.get("status")),
            "Last_Validated": now if item.get("validation_performed") else None,
            "Quality_Report_ID": f"DQ_{data_id}" if data_id else None,
            "Notes": _catalog_notes(item),
        }
        write_named_row(
            sheet,
            DATA_CATALOG_HEADER_ROW,
            DATA_CATALOG_FIRST_ROW + offset,
            values,
            required=("Data_ID",),
        )


def write_data_quality_report(
    sheet: Worksheet,
    manifest: dict[str, Any],
    run_id: str,
    engine_version: str,
    project_root: Path,
    now: datetime,
    manifest_path: Path | None = None,
) -> None:
    files = list(manifest.get("files") or [])
    _require_capacity("Data_Quality_Report", len(files), DATA_QUALITY_FIRST_ROW, DATA_QUALITY_LAST_ROW)
    _clear_rows(sheet, DATA_QUALITY_FIRST_ROW, DATA_QUALITY_LAST_ROW)
    _report_choices(sheet, "D4:D35", "Pending,Pass,Review,Fail,Validated,Failed,Headers Checked,Partially Checked")
    _report_choices(sheet, "J4:J35", "Yes,No,Consistent,Inconclusive,Partial scan,Not Checked", replaced="J4:K35")
    _report_choices(sheet, "K4:K35", "Yes,No,Not Checked")
    headers = table_headers(sheet, DATA_QUALITY_HEADER_ROW)

    report_path = relative_to_root(manifest_path, project_root) if manifest_path else (
        f"archive/backtest_runs/{run_id}/data_manifest.json" if run_id else ""
    )
    for offset, item in enumerate(files):
        data_id = item.get("data_id")
        diagnostics = item.get("diagnostics") or {}
        gaps = diagnostics.get("missing_bar_check") or {}
        spread = diagnostics.get("spread_check") or {}
        maximum = spread.get("maximum") or {}
        timezone_check = diagnostics.get("timezone_check") or {}
        values = {
            "Report_ID": f"DQ_{data_id}" if data_id else None,
            "Data_ID": data_id,
            "Check_DateTime": now if item.get("validation_performed") else None,
            "Check_Status": item.get("report_status", item.get("status")),
            "Missing_Bars": gaps.get("missing_bars") if gaps.get("missing_bars") is not None
            else gaps.get("status", "Not Checked"),
            "Duplicate_Timestamps": _checked_count(item, "duplicate_timestamps"),
            "Bad_Bid_Ask": _checked_count(item, "bad_bid_ask_rows"),
            "Negative_Spread": _checked_count(item, "negative_spread_rows"),
            "Max_Spread": maximum.get("pips") if maximum.get("pips") is not None
            else spread.get("status", "Not Checked"),
            "Timezone_OK": timezone_check.get("status", "Not Checked"),
            "Date_Range_OK": ("Yes" if item["date_range_ok"] else "No")
            if item.get("date_range_ok") is not None else "Not Checked",
            "Rows_Checked": _rows_checked(item),
            "Engine_Version": engine_version,
            "Report_Path": relative_to_root(project_root / report_path, project_root) if report_path else "",
            "Notes": _quality_notes(item),
        }
        write_named_row(
            sheet,
            DATA_QUALITY_HEADER_ROW,
            DATA_QUALITY_FIRST_ROW + offset,
            values,
            required=("Report_ID", "Data_ID"),
        )
        if "Max_Spread" in headers:
            sheet.cell(DATA_QUALITY_FIRST_ROW + offset, headers["Max_Spread"]).number_format = '0.00" pips"'


def _report_choices(sheet: Worksheet, cell_range: str, choices: str, replaced: str | None = None) -> None:
    for validation in list(sheet.data_validations.dataValidation):
        if str(validation.sqref) in {cell_range, replaced}:
            sheet.data_validations.dataValidation.remove(validation)
    validation = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
    validation.add(cell_range)
    sheet.add_data_validation(validation)


def append_analysis_run_id(sheet: Worksheet, run_id: str) -> None:
    headers = table_headers(sheet, ANALYSIS_HEADER_ROW)
    run_col = headers.get("Run_ID")
    if run_col is None:
        raise ExcelWriterError(f"{sheet.title} is missing required column 'Run_ID'.")

    for row_idx in range(ANALYSIS_FIRST_ROW, ANALYSIS_LAST_ROW + 1):
        if sheet.cell(row=row_idx, column=run_col).value == run_id:
            return

    for row_idx in range(ANALYSIS_FIRST_ROW, ANALYSIS_LAST_ROW + 1):
        if sheet.cell(row=row_idx, column=run_col).value in (None, ""):
            sheet.cell(row=row_idx, column=run_col, value=run_id)
            return

    raise ExcelWriterError(f"{sheet.title} has no empty analysis rows left.")


def write_analysis_run_summary(
    sheet: Worksheet,
    planned: Any,
    selected: Any,
    config: dict[str, Any],
    engine_version: str,
    data_version: str,
    summary: dict[str, Any] | None,
    activate: bool,
    official_oos: bool,
) -> None:
    row_idx = _analysis_row_for_run(sheet, planned.run_id)
    is_active = activate and (planned.sample_type == "IS" or official_oos)
    values = {
        "Run_ID": planned.run_id,
        "Window_ID": selected.window_id,
        "Sample_Type": planned.sample_type,
        "Active_For_Analysis": "Yes" if is_active else "No",
        "Period_Start": planned.period_start,
        "Period_End": planned.period_end,
        "Rule_Version": config.get("Rule_Version"),
        "Engine_Version": engine_version,
        "Data_Version": data_version,
        "Trades": _summary_value(summary, "trades"),
        "Wins": _summary_value(summary, "wins"),
        "Losses": _summary_value(summary, "losses"),
        "Win_Rate": _summary_value(summary, "win_rate"),
        "Net_R": _summary_value(summary, "net_r"),
        "Avg_R": _summary_value(summary, "avg_r"),
        "Trade_Based_Sharpe": _summary_value(summary, "trade_based_sharpe", "trade_sharpe"),
        "Annualised_Trade_Sharpe": _summary_value(
            summary,
            "annualised_trade_sharpe",
            "annualized_trade_sharpe",
            "annualised_sharpe",
        ),
        "Daily_Sharpe": _summary_value(summary, "daily_sharpe"),
        "Annualised_Daily_Sharpe": _summary_value(
            summary,
            "annualised_daily_sharpe",
            "annualized_daily_sharpe",
        ),
        "Max_Drawdown_R": _summary_value(summary, "max_drawdown_r"),
        "Profit_Factor": _summary_value(summary, "profit_factor"),
        "Pass_Fail": "Pending",
    }
    write_named_row(sheet, ANALYSIS_HEADER_ROW, row_idx, values, required=("Run_ID",))


def empty_year_analysis_rows(year_analysis: Worksheet, count: int) -> list[int]:
    return empty_rows(
        sheet=year_analysis,
        id_column="Year_Record_ID",
        header_row=YEAR_ANALYSIS_HEADER_ROW,
        first_row=YEAR_ANALYSIS_FIRST_ROW,
        last_row=YEAR_ANALYSIS_LAST_ROW,
        count=count,
        label="Year_Analysis",
    )


def empty_is_run_review_rows(sheet: Worksheet, count: int) -> list[int]:
    return empty_rows(
        sheet=sheet,
        id_column="Review_Record_ID",
        header_row=IS_REVIEW_HEADER_ROW,
        first_row=IS_REVIEW_FIRST_ROW,
        last_row=IS_REVIEW_LAST_ROW,
        count=count,
        label="IS_Run_Review",
    )


def empty_parameter_snapshot_rows(sheet: Worksheet, count: int) -> list[int]:
    return empty_rows(
        sheet=sheet,
        id_column="Snapshot_ID",
        header_row=PARAMETER_SNAPSHOT_HEADER_ROW,
        first_row=PARAMETER_SNAPSHOT_FIRST_ROW,
        last_row=PARAMETER_SNAPSHOT_LAST_ROW,
        count=count,
        label="Parameter_Snapshot",
    )


def empty_is_run_comparison_rows(sheet: Worksheet, count: int) -> list[int]:
    return empty_rows(
        sheet=sheet,
        id_column="Comparison_Record_ID",
        header_row=IS_COMPARISON_HEADER_ROW,
        first_row=IS_COMPARISON_FIRST_ROW,
        last_row=IS_COMPARISON_LAST_ROW,
        count=count,
        label="IS_Run_Comparison",
    )


def empty_equity_curve_rows(sheet: Worksheet, count: int) -> list[int]:
    return empty_rows(
        sheet=sheet,
        id_column="Equity_Record_ID",
        header_row=EQUITY_CURVE_HEADER_ROW,
        first_row=EQUITY_CURVE_FIRST_ROW,
        last_row=EQUITY_CURVE_LAST_ROW,
        count=count,
        label="Equity_Curve",
    )


def empty_signal_audit_rows(sheet: Worksheet, count: int) -> list[int]:
    return empty_rows(
        sheet=sheet,
        id_column="Audit_Record_ID",
        header_row=SIGNAL_AUDIT_HEADER_ROW,
        first_row=SIGNAL_AUDIT_FIRST_ROW,
        last_row=SIGNAL_AUDIT_LAST_ROW,
        count=count,
        label="Signal_Audit",
    )


def empty_oos_result_review_rows(sheet: Worksheet, count: int) -> list[int]:
    return empty_rows(
        sheet=sheet,
        id_column="Review_Record_ID",
        header_row=OOS_RESULT_REVIEW_HEADER_ROW,
        first_row=OOS_RESULT_REVIEW_FIRST_ROW,
        last_row=OOS_RESULT_REVIEW_LAST_ROW,
        count=count,
        label="OOS_Result_Review",
    )


def append_year_analysis_rows(
    sheet: Worksheet,
    rows: list[int],
    planned: Any,
    selected: Any,
    config: dict[str, Any],
    engine_version: str,
    data_version: str,
    yearly_summary: dict[int, dict[str, Any]] | dict[str, dict[str, Any]] | None,
    activate: bool,
    official_oos: bool,
) -> None:
    yearly_items = _sorted_yearly_items(yearly_summary)
    if len(rows) != len(yearly_items):
        raise ExcelWriterError(
            f"Year_Analysis row count mismatch for {planned.run_id}: "
            f"{len(rows)} rows for {len(yearly_items)} years."
        )

    is_official_oos = planned.sample_type == "OOS" and official_oos
    is_active = activate and (planned.sample_type == "IS" or is_official_oos)
    for row_idx, (year, summary) in zip(rows, yearly_items):
        values = {
            "Year_Record_ID": f"{planned.run_id}_{year}",
            "Run_ID": planned.run_id,
            "Window_ID": selected.window_id,
            "Sample_Type": planned.sample_type,
            "Year": year,
            "Active_For_Analysis": "Yes" if is_active else "No",
            "Official_OOS_Run": "Yes" if is_official_oos else "No",
            "Rule_Version": config.get("Rule_Version"),
            "Engine_Version": engine_version,
            "Data_Version": data_version,
            "Trades": _summary_value(summary, "trades"),
            "Wins": _summary_value(summary, "wins"),
            "Losses": _summary_value(summary, "losses"),
            "Win_Rate": _summary_value(summary, "win_rate"),
            "Net_R": _summary_value(summary, "net_r"),
            "Avg_R": _summary_value(summary, "avg_r"),
            "Trade_Based_Sharpe": _summary_value(summary, "trade_based_sharpe", "trade_sharpe"),
            "Annualised_Trade_Sharpe": _summary_value(
                summary,
                "annualised_trade_sharpe",
                "annualized_trade_sharpe",
                "annualised_sharpe",
            ),
            "Daily_Sharpe": _summary_value(summary, "daily_sharpe"),
            "Annualised_Daily_Sharpe": _summary_value(
                summary,
                "annualised_daily_sharpe",
                "annualized_daily_sharpe",
            ),
            "Max_Drawdown_R": _summary_value(summary, "max_drawdown_r"),
            "Profit_Factor": _summary_value(summary, "profit_factor"),
            "Best_R": _summary_value(summary, "best_r"),
            "Worst_R": _summary_value(summary, "worst_r"),
            "News_Blocked_Signals": _summary_value(summary, "news_blocked_signals"),
            "Valid_Signals": _summary_value(summary, "valid_signals"),
            "Invalid_Signals": _summary_value(summary, "invalid_signals"),
            "No_Fill": _summary_value(summary, "no_fill"),
            "Blocked": _summary_value(summary, "blocked"),
            "Forced_Closes": _summary_value(summary, "forced"),
            "Notes": "",
        }
        write_named_row(
            sheet,
            YEAR_ANALYSIS_HEADER_ROW,
            row_idx,
            values,
            required=("Year_Record_ID", "Run_ID", "Year"),
        )
        apply_table_body_row_theme(sheet, YEAR_ANALYSIS_HEADERS, row_idx)


def append_parameter_snapshot_row(
    sheet: Worksheet,
    row_idx: int,
    planned: Any,
    selected: Any,
    config: dict[str, Any],
    rule_hash: str,
    engine_version: str,
    engine_hash: str,
    data_version: str,
    strategy_params_hash: str,
    key_params: str,
    project_root: Path,
    now: datetime,
) -> None:
    values = {
        "Snapshot_ID": f"{planned.run_id}_PARAMS",
        "Run_ID": planned.run_id,
        "Window_ID": selected.window_id,
        "Sample_Type": planned.sample_type,
        "Run_DateTime": now,
        "Rule_Version": config.get("Rule_Version"),
        "Rule_Set_ID": config.get("Rule_Set_ID"),
        "Rule_Config_Hash": rule_hash,
        "Engine_Version": engine_version,
        "Engine_Config_Hash": engine_hash,
        "Data_Version": data_version,
        "Key_Params": key_params,
        "Strategy_Params_Hash": strategy_params_hash,
        "Strategy_Params_Path": relative_to_root(planned.archive_dir / "strategy_params.json", project_root),
        "Rule_Config_Path": relative_to_root(planned.archive_dir / "rule_config.json", project_root),
        "Engine_Config_Path": relative_to_root(planned.archive_dir / "engine_config.json", project_root),
        "Archive_Path": relative_to_root(planned.archive_dir, project_root),
        "Notes": "",
    }
    write_named_row(
        sheet,
        PARAMETER_SNAPSHOT_HEADER_ROW,
        row_idx,
        values,
        required=("Snapshot_ID", "Run_ID"),
    )
    apply_table_body_row_theme(sheet, PARAMETER_SNAPSHOT_HEADERS, row_idx)


def append_is_run_review_row(
    sheet: Worksheet,
    row_idx: int,
    planned: Any,
    selected: Any,
    config: dict[str, Any],
    rule_hash: str,
    engine_hash: str,
    summary: dict[str, Any] | None,
    yearly_summary: dict[int, dict[str, Any]] | dict[str, dict[str, Any]] | None,
    now: datetime,
) -> None:
    values = {
        "Review_Record_ID": f"{planned.run_id}_REVIEW",
        "Window_ID": selected.window_id,
        "Candidate_IS_Run_ID": planned.run_id,
        "Run_DateTime": now,
        "IS_Period": _period_label(planned),
        "Keep_Reject": "Review",
        "Reason": "",
        "Selected_For_OOS": "No",
        "Locked_At": "",
        "Locked_By": "",
        "Rule_Version": config.get("Rule_Version"),
        "Rule_Config_Hash": rule_hash,
        "Engine_Config_Hash": engine_hash,
        "Trades": _summary_value(summary, "trades"),
        "Net_R": _summary_value(summary, "net_r"),
        "Max_Drawdown_R": _summary_value(summary, "max_drawdown_r"),
        "Trade_Based_Sharpe": _summary_value(summary, "trade_based_sharpe", "trade_sharpe"),
        "Annualised_Trade_Sharpe": _summary_value(
            summary,
            "annualised_trade_sharpe",
            "annualized_trade_sharpe",
            "annualised_sharpe",
        ),
        "Daily_Sharpe": _summary_value(summary, "daily_sharpe"),
        "Annualised_Daily_Sharpe": _summary_value(
            summary,
            "annualised_daily_sharpe",
            "annualized_daily_sharpe",
        ),
        "Profit_Factor": _summary_value(summary, "profit_factor"),
        "Year_Consistency": _year_consistency(yearly_summary),
        "Review_Status": "Pending",
        "Notes": "",
    }
    write_named_row(
        sheet,
        IS_REVIEW_HEADER_ROW,
        row_idx,
        values,
        required=("Review_Record_ID", "Candidate_IS_Run_ID"),
    )
    apply_table_body_row_theme(sheet, IS_RUN_REVIEW_HEADERS, row_idx)


def append_is_run_comparison_row(
    sheet: Worksheet,
    row_idx: int,
    planned: Any,
    selected: Any,
    config: dict[str, Any],
    rule_hash: str,
    engine_hash: str,
    data_version: str,
    summary: dict[str, Any] | None,
    yearly_summary: dict[int, dict[str, Any]] | dict[str, dict[str, Any]] | None,
    now: datetime,
    activate: bool,
) -> None:
    values = {
        "Comparison_Record_ID": f"{planned.run_id}_COMPARE",
        "Window_ID": selected.window_id,
        "Run_ID": planned.run_id,
        "Run_DateTime": now,
        "IS_Period": _period_label(planned),
        "Active_For_Analysis": "Yes" if activate else "No",
        "Keep_Reject": _is_review_lookup_formula(f"C{row_idx}", "F", "Review"),
        "Selected_For_OOS": _is_review_lookup_formula(f"C{row_idx}", "H", "No"),
        "Rule_Version": config.get("Rule_Version"),
        "Rule_Config_Hash": rule_hash,
        "Engine_Config_Hash": engine_hash,
        "Data_Version": data_version,
        "Trades": _summary_value(summary, "trades"),
        "Wins": _summary_value(summary, "wins"),
        "Losses": _summary_value(summary, "losses"),
        "Win_Rate": _summary_value(summary, "win_rate"),
        "Net_R": _summary_value(summary, "net_r"),
        "Max_Drawdown_R": _summary_value(summary, "max_drawdown_r"),
        "Trade_Based_Sharpe": _summary_value(summary, "trade_based_sharpe", "trade_sharpe"),
        "Annualised_Trade_Sharpe": _summary_value(
            summary,
            "annualised_trade_sharpe",
            "annualized_trade_sharpe",
            "annualised_sharpe",
        ),
        "Daily_Sharpe": _summary_value(summary, "daily_sharpe"),
        "Annualised_Daily_Sharpe": _summary_value(
            summary,
            "annualised_daily_sharpe",
            "annualized_daily_sharpe",
        ),
        "Profit_Factor": _summary_value(summary, "profit_factor"),
        "Best_R": _summary_value(summary, "best_r"),
        "Worst_R": _summary_value(summary, "worst_r"),
        "News_Blocked_Signals": _summary_value(summary, "news_blocked_signals"),
        "Valid_Signals": _summary_value(summary, "valid_signals"),
        "Invalid_Signals": _summary_value(summary, "invalid_signals"),
        "Year_Consistency": _year_consistency(yearly_summary),
        "Review_Notes": _is_review_lookup_formula(f"C{row_idx}", "G", ""),
    }
    write_named_row(
        sheet,
        IS_COMPARISON_HEADER_ROW,
        row_idx,
        values,
        required=("Comparison_Record_ID", "Run_ID"),
    )
    apply_table_body_row_theme(sheet, IS_RUN_COMPARISON_HEADERS, row_idx)


def append_equity_curve_rows(
    sheet: Worksheet,
    rows: list[int],
    equity_curve: list[dict[str, Any]],
    planned: Any,
    selected: Any,
) -> None:
    if len(rows) != len(equity_curve):
        raise ExcelWriterError(
            f"Equity_Curve row count mismatch for {planned.run_id}: "
            f"{len(rows)} rows for {len(equity_curve)} equity points."
        )
    for row_idx, point in zip(rows, equity_curve):
        trade_number = point.get("trade_number")
        values = {
            "Equity_Record_ID": f"{planned.run_id}_EQ{int(trade_number or 0):04d}",
            "Run_ID": planned.run_id,
            "Window_ID": selected.window_id,
            "Sample_Type": planned.sample_type,
            "Trade_Number": trade_number,
            "Trade_ID": f"{planned.run_id}_T{int(trade_number or 0):04d}",
            "Signal_ID": point.get("signal_id"),
            "Entry_DateTime": point.get("entry_time"),
            "Exit_DateTime": point.get("exit_time"),
            "Direction": point.get("direction"),
            "Setup_Tag": point.get("setup_type"),
            "Outcome": point.get("outcome"),
            "Exit_Reason": point.get("exit_reason"),
            "R_Multiple": point.get("r_multiple"),
            "PnL_Pct": point.get("pnl_pct"),
            "Cumulative_R": point.get("cumulative_r"),
            "Equity_Peak_R": point.get("equity_peak_r"),
            "Drawdown_R": point.get("drawdown_r"),
            "Cumulative_Return_Pct": point.get("cumulative_return_pct"),
            "Notes": "",
        }
        write_named_row(
            sheet,
            EQUITY_CURVE_HEADER_ROW,
            row_idx,
            values,
            required=("Equity_Record_ID", "Run_ID"),
        )
        apply_table_body_row_theme(sheet, EQUITY_CURVE_HEADERS, row_idx)


def append_signal_audit_rows(
    sheet: Worksheet,
    rows: list[int],
    signal_audit: list[dict[str, Any]],
    planned: Any,
    selected: Any,
) -> None:
    if len(rows) != len(signal_audit):
        raise ExcelWriterError(
            f"Signal_Audit row count mismatch for {planned.run_id}: "
            f"{len(rows)} rows for {len(signal_audit)} signal audit entries."
        )
    for row_idx, signal in zip(rows, signal_audit):
        signal_id = signal.get("signal_id") or f"SIG_{row_idx:05d}"
        values = {
            "Audit_Record_ID": f"{planned.run_id}_{signal_id}",
            "Run_ID": planned.run_id,
            "Window_ID": selected.window_id,
            "Sample_Type": planned.sample_type,
            "Signal_ID": signal_id,
            "Signal_DateTime": signal.get("signal_time"),
            "Direction": signal.get("direction"),
            "Setup_Type": signal.get("setup_type"),
            "Status": signal.get("status"),
            "Reject_Category": signal.get("reject_category"),
            "Reason": signal.get("reason"),
            "Obstacle_Hit": _yes_no(signal.get("obstacle_hit")),
            "News_Blocked": _yes_no(signal.get("news_blocked")),
            "News_Event_IDs": signal.get("news_event_ids"),
            "News_Events": signal.get("news_events"),
            "Run_Count": signal.get("run_count"),
            "Largest_Body_Pips": signal.get("largest_body_pips"),
            "Entry_Frozen": _yes_no(signal.get("entry_frozen")),
            "Entry": signal.get("entry"),
            "Stop_Loss": signal.get("stop_loss"),
            "Take_Profit": signal.get("take_profit"),
            "Source_Row": signal.get("source_row"),
            "Run_Start": signal.get("run_start"),
            "Run_End": signal.get("run_end"),
            "Signal_Text": signal.get("signal_text"),
            "Extension_Freeze": _yes_no(signal.get("extension_freeze")),
            "Extension_Freeze_Reason": signal.get("extension_freeze_reason"),
            "Extension_Freeze_Touch_Time": signal.get("extension_freeze_touch_time"),
            "Extension_Freeze_Candidate_Time": signal.get("extension_freeze_candidate_time"),
            "Extension_Freeze_Line_Version": signal.get("extension_freeze_line_version"),
            "Context_Filter_Blocked": _yes_no(signal.get("context_filter_blocked")),
            "Context_Filter_Reason": signal.get("context_filter_reason"),
            "Adverse_Asia_Reversal_Sell_Filter_Triggered": _yes_no(
                signal.get("adverse_asia_reversal_sell_filter_triggered")
            ),
            "Active_Price_Cluster_Stop_Filter_Triggered": _yes_no(
                signal.get("active_price_cluster_stop_filter_triggered")
            ),
            "Adverse_Order_Block_Stop_Filter_Triggered": _yes_no(
                signal.get("adverse_order_block_stop_filter_triggered")
            ),
            "Pre_Run_Structure_Direction": signal.get("pre_run_structure_direction"),
            "Previous_Day_Signed_PFE_Trade_Direction": signal.get(
                "previous_day_signed_pfe_trade_direction"
            ),
            "Pre_Run_Asia_Signed_PFE_Trade_Direction": signal.get(
                "pre_run_asia_signed_pfe_trade_direction"
            ),
            "Adverse_Asia_Reversal_Sell": _yes_no(
                signal.get("adverse_asia_reversal_sell")
            ),
            "Active_Cluster_Present_Ahead": _yes_no(
                signal.get("active_cluster_present_ahead")
            ),
            "Active_Cluster_Count_Ahead": signal.get("active_cluster_count_ahead"),
            "Active_Terminal_Cluster_ID": signal.get("active_terminal_cluster_id"),
            "Active_Terminal_Cluster_Low": signal.get("active_terminal_cluster_low"),
            "Active_Terminal_Cluster_High": signal.get("active_terminal_cluster_high"),
            "Active_Terminal_Cluster_Source_Move_ID": signal.get(
                "active_terminal_cluster_source_move_id"
            ),
            "Active_Terminal_Cluster_Source_Move_Start_Time": signal.get(
                "active_terminal_cluster_source_move_start_time"
            ),
            "Active_Terminal_Cluster_Source_Move_End_Time": signal.get(
                "active_terminal_cluster_source_move_end_time"
            ),
            "Active_Cluster_Interaction": signal.get("active_cluster_interaction"),
            "Run_Stopped_At_Active_Cluster": _yes_no(
                signal.get("run_stopped_at_active_cluster")
            ),
            "Run_Cleared_Active_Cluster_Count": signal.get(
                "run_cleared_active_cluster_count"
            ),
            "Active_Order_Block_Count_Ahead": signal.get(
                "active_order_block_count_ahead"
            ),
            "Terminal_Order_Block_ID": signal.get("terminal_order_block_id"),
            "Terminal_Order_Block_Low": signal.get("terminal_order_block_low"),
            "Terminal_Order_Block_High": signal.get("terminal_order_block_high"),
            "Terminal_Order_Block_Source_Move_ID": signal.get(
                "terminal_order_block_source_move_id"
            ),
            "Terminal_Order_Block_Source_Move_Start_Time": signal.get(
                "terminal_order_block_source_move_start_time"
            ),
            "Terminal_Order_Block_Source_Move_End_Time": signal.get(
                "terminal_order_block_source_move_end_time"
            ),
            "Order_Block_Broad_Interaction": signal.get(
                "order_block_broad_interaction"
            ),
            "Order_Block_Reached_Without_Clean_Close": _yes_no(
                signal.get("order_block_reached_without_clean_close")
            ),
            "Run_Cleared_Order_Block_Count": signal.get(
                "run_cleared_order_block_count"
            ),
            "Adverse_Order_Block_Stop": _yes_no(
                signal.get("adverse_order_block_stop")
            ),
        }
        write_named_row(
            sheet,
            SIGNAL_AUDIT_HEADER_ROW,
            row_idx,
            values,
            required=("Audit_Record_ID", "Run_ID", "Signal_ID"),
        )
        apply_table_body_row_theme(sheet, SIGNAL_AUDIT_HEADERS, row_idx)


def append_oos_result_review_row(
    sheet: Worksheet,
    row_idx: int,
    planned: Any,
    selected: Any,
    config: dict[str, Any],
    rule_hash: str,
    engine_hash: str,
    linked_is_run_id: str,
    summary: dict[str, Any] | None,
    yearly_summary: dict[int, dict[str, Any]] | dict[str, dict[str, Any]] | None,
    now: datetime,
    official_oos: bool,
) -> None:
    values = {
        "Review_Record_ID": f"{planned.run_id}_OOS_REVIEW",
        "Window_ID": selected.window_id,
        "OOS_Run_ID": planned.run_id,
        "Linked_IS_Run_ID": linked_is_run_id,
        "Run_DateTime": now,
        "OOS_Period": _period_label(planned),
        "OOS_Year": selected.oos_year,
        "Official_OOS_Run": "Yes" if official_oos else "No",
        "OOS_Locked": "Yes" if official_oos else "No",
        "Review_Status": "Pending",
        "Pass_Fail": "Review",
        "Continue_To_Next_Window": "",
        "Decision_Reason": "",
        "Reviewed_At": "",
        "Reviewed_By": "",
        "Rule_Version": config.get("Rule_Version"),
        "Rule_Config_Hash": rule_hash,
        "Engine_Config_Hash": engine_hash,
        "Trades": _summary_value(summary, "trades"),
        "Net_R": _summary_value(summary, "net_r"),
        "Max_Drawdown_R": _summary_value(summary, "max_drawdown_r"),
        "Trade_Based_Sharpe": _summary_value(summary, "trade_based_sharpe", "trade_sharpe"),
        "Annualised_Trade_Sharpe": _summary_value(
            summary,
            "annualised_trade_sharpe",
            "annualized_trade_sharpe",
            "annualised_sharpe",
        ),
        "Daily_Sharpe": _summary_value(summary, "daily_sharpe"),
        "Annualised_Daily_Sharpe": _summary_value(
            summary,
            "annualised_daily_sharpe",
            "annualized_daily_sharpe",
        ),
        "Profit_Factor": _summary_value(summary, "profit_factor"),
        "Year_Consistency": _year_consistency(yearly_summary),
        "Notes": "",
    }
    write_named_row(
        sheet,
        OOS_RESULT_REVIEW_HEADER_ROW,
        row_idx,
        values,
        required=("Review_Record_ID", "OOS_Run_ID"),
    )
    apply_table_body_row_theme(sheet, OOS_RESULT_REVIEW_HEADERS, row_idx)


def empty_rows(
    sheet: Worksheet,
    id_column: str,
    header_row: int,
    first_row: int,
    last_row: int,
    count: int,
    label: str,
) -> list[int]:
    if count <= 0:
        return []

    headers = table_headers(sheet, header_row)
    id_col = headers.get(id_column)
    if id_col is None:
        raise ExcelWriterError(f"{label} is missing required column {id_column!r}.")

    rows: list[int] = []
    for row_idx in range(first_row, last_row + 1):
        if sheet.cell(row=row_idx, column=id_col).value in (None, ""):
            rows.append(row_idx)
            if len(rows) == count:
                return rows
    raise ExcelWriterError(f"{label} has no empty rows left. Expand the workbook table.")


def append_run_log_row(
    sheet: Worksheet,
    row_idx: int,
    planned: Any,
    selected: Any,
    config: dict[str, Any],
    rule_hash: str,
    engine_version: str,
    engine_hash: str,
    data_version: str,
    project_root: Path,
    now: datetime,
    activate: bool,
    official_oos: bool,
    run_message: str,
) -> None:
    is_oos = planned.sample_type == "OOS"
    values = {
        "Run_ID": planned.run_id,
        "Run_DateTime": now,
        "Window_ID": selected.window_id,
        "Sample_Type": planned.sample_type,
        "Rule_Version": config.get("Rule_Version"),
        "Rule_Set_ID": config.get("Rule_Set_ID"),
        "Rule_Config_Hash": rule_hash,
        "Engine_Version": engine_version,
        "Engine_Config_Hash": engine_hash,
        "Data_Version": data_version,
        "Run_Status": "Complete",
        "Active_For_Analysis": "Yes" if activate else "No",
        "Official_OOS_Run": "Yes" if is_oos and official_oos else "No",
        "Snapshot_Created": "Yes",
        "Snapshot_Path": relative_to_root(planned.workbook_after_path, project_root),
        "Workbook_Before_Path": relative_to_root(planned.workbook_before_path, project_root),
        "Period_Start": planned.period_start,
        "Period_End": planned.period_end,
        "Run_Message": run_message,
        "Notes": "",
        "Archive_Notes": relative_to_root(planned.archive_dir, project_root),
    }
    write_named_row(sheet, RUN_LOG_HEADER_ROW, row_idx, values, required=("Run_ID",))


def append_trade_log_rows(
    sheet: Worksheet,
    rows: list[int],
    trades: list[dict[str, Any]],
    planned: Any,
    selected: Any,
    config: dict[str, Any],
) -> None:
    for trade_number, (row_idx, trade) in enumerate(zip(rows, trades), start=1):
        values = trade_log_values(
            trade=trade,
            trade_id=f"{planned.run_id}_T{trade_number:04d}",
            planned=planned,
            selected=selected,
            config=config,
        )
        write_named_row(sheet, TRADE_LOG_HEADER_ROW, row_idx, values, required=("Trade_ID", "Run_ID"))


def trade_log_values(
    trade: dict[str, Any],
    trade_id: str,
    planned: Any,
    selected: Any,
    config: dict[str, Any],
) -> dict[str, Any]:
    pips = trade.get("pips")
    result = normalize_result(trade.get("outcome"))
    review = "Review" if result == "Open" else ""
    notes = trade.get("blocked_reason") or trade.get("exit_reason") or trade.get("reason") or ""
    return {
        "Trade_ID": trade_id,
        "Run_ID": planned.run_id,
        "Window_ID": selected.window_id,
        "Sample_Type": planned.sample_type,
        "Symbol": "GBP_USD",
        "Direction": "Long" if trade.get("direction") == "Buy" else "Short",
        "Setup_Tag": trade.get("setup_type"),
        "Signal_ID": trade.get("signal_id"),
        "Signal_DateTime": trade.get("signal_time"),
        "Entry_DateTime": trade.get("entry_time"),
        "Exit_DateTime": trade.get("exit_time"),
        "Entry_Price": trade.get("entry_price"),
        "Exit_Price": trade.get("exit_price"),
        "Stop_Price": trade.get("stop_price"),
        "Target_Price": trade.get("target_price"),
        "Risk_Points": trade.get("risk_pips"),
        "Gross_PnL_Points": pips,
        "Fees_Points": 0,
        "Net_PnL_Points": pips,
        "R_Multiple": trade.get("r_multiple"),
        "Result": result,
        "Entry_Timeframe": "15m",
        "Execution_Timeframe": "1m",
        "Spread_At_Entry": trade.get("spread_at_entry"),
        "Spread_At_Exit": None,
        "Rule_Version": config.get("Rule_Version"),
        "Data_Quality_Flag": "OK",
        "Review_Flag": review,
        "Notes": notes,
    }


def write_named_row(
    sheet: Worksheet,
    header_row: int,
    row_idx: int,
    values: dict[str, Any],
    required: tuple[str, ...] = (),
) -> None:
    headers = table_headers(sheet, header_row)
    missing = [name for name in required if name not in headers]
    if missing:
        raise ExcelWriterError(
            f"{sheet.title} is missing required column(s): {', '.join(missing)}."
        )

    for name, value in values.items():
        col_idx = headers.get(name)
        if col_idx is not None:
            sheet.cell(row=row_idx, column=col_idx, value=value)


def _clear_rows(sheet: Worksheet, first_row: int, last_row: int) -> None:
    for row_idx in range(first_row, last_row + 1):
        for col_idx in range(1, sheet.max_column + 1):
            sheet.cell(row=row_idx, column=col_idx).value = None


def _copy_column_style(sheet: Worksheet, source_col: int, target_col: int) -> None:
    for row_idx in range(1, sheet.max_row + 1):
        source = sheet.cell(row=row_idx, column=source_col)
        target = sheet.cell(row=row_idx, column=target_col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)


def _write_wf_metric_formula(
    sheet: Worksheet,
    row_idx: int,
    target_col: int,
    analysis_sheet: str,
    metric_header: str,
    run_ref_col: str,
) -> None:
    metric_col = _analysis_column_letter(analysis_sheet, metric_header)
    active_col = _analysis_column_letter(analysis_sheet, "Active_For_Analysis")
    run_col = _analysis_column_letter(analysis_sheet, "Run_ID")
    sheet.cell(
        row=row_idx,
        column=target_col,
        value=(
            f'=IF({run_ref_col}{row_idx}="",0,'
            f'SUMIFS(\'{analysis_sheet}\'!${metric_col}$4:${metric_col}$35,'
            f'\'{analysis_sheet}\'!${run_col}$4:${run_col}$35,{run_ref_col}{row_idx},'
            f'\'{analysis_sheet}\'!${active_col}$4:${active_col}$35,"Yes"))'
        ),
    )


def _analysis_column_letter(sheet_name: str, header: str) -> str:
    default_maps = {
        "IS_Analysis": {
            "Run_ID": "A",
            "Active_For_Analysis": "D",
            "Trades": "J",
            "Wins": "K",
            "Losses": "L",
            "Win_Rate": "M",
            "Net_R": "N",
            "Trade_Based_Sharpe": "P",
            "Annualised_Trade_Sharpe": "Q",
            "Daily_Sharpe": "R",
            "Annualised_Daily_Sharpe": "S",
        },
        "OOS_Analysis": {
            "Run_ID": "A",
            "Active_For_Analysis": "D",
            "Trades": "J",
            "Wins": "K",
            "Losses": "L",
            "Win_Rate": "M",
            "Net_R": "N",
            "Trade_Based_Sharpe": "P",
            "Annualised_Trade_Sharpe": "Q",
            "Daily_Sharpe": "R",
            "Annualised_Daily_Sharpe": "S",
        },
    }
    try:
        return default_maps[sheet_name][header]
    except KeyError as exc:
        raise ExcelWriterError(f"No default column mapping for {sheet_name}.{header}.") from exc


def _analysis_row_for_run(sheet: Worksheet, run_id: str) -> int:
    headers = table_headers(sheet, ANALYSIS_HEADER_ROW)
    run_col = headers.get("Run_ID")
    if run_col is None:
        raise ExcelWriterError(f"{sheet.title} is missing required column 'Run_ID'.")

    for row_idx in range(ANALYSIS_FIRST_ROW, ANALYSIS_LAST_ROW + 1):
        if sheet.cell(row=row_idx, column=run_col).value == run_id:
            return row_idx

    for row_idx in range(ANALYSIS_FIRST_ROW, ANALYSIS_LAST_ROW + 1):
        if sheet.cell(row=row_idx, column=run_col).value in (None, ""):
            return row_idx

    raise ExcelWriterError(f"{sheet.title} has no empty analysis rows left.")


def _summary_value(summary: dict[str, Any] | None, *keys: str) -> Any:
    if not summary:
        return 0
    for key in keys:
        value = summary.get(key)
        if value not in (None, ""):
            return value
    return 0


def _yes_no(value: Any) -> str:
    return "Yes" if bool(value) else "No"


def _period_label(planned: Any) -> str:
    start = _format_date(planned.period_start)
    end = _format_date(planned.period_end)
    return f"{start} to {end}"


def _format_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _year_consistency(
    yearly_summary: dict[int, dict[str, Any]] | dict[str, dict[str, Any]] | None,
) -> float:
    yearly_items = _sorted_yearly_items(yearly_summary)
    if not yearly_items:
        return 0.0
    positive_years = sum(1 for _, summary in yearly_items if _summary_value(summary, "net_r") > 0)
    return positive_years / len(yearly_items)


def _is_review_lookup_formula(run_id_ref: str, return_col: str, default: str) -> str:
    escaped_default = default.replace('"', '""')
    return (
        f'=IFERROR(INDEX(\'IS_Run_Review\'!${return_col}${IS_REVIEW_FIRST_ROW}:'
        f'${return_col}${IS_REVIEW_LAST_ROW},'
        f'MATCH({run_id_ref},\'IS_Run_Review\'!$C${IS_REVIEW_FIRST_ROW}:'
        f'$C${IS_REVIEW_LAST_ROW},0)),"{escaped_default}")'
    )


def _sorted_yearly_items(
    yearly_summary: dict[int, dict[str, Any]] | dict[str, dict[str, Any]] | None,
) -> list[tuple[int, dict[str, Any]]]:
    if not yearly_summary:
        return []
    output: list[tuple[int, dict[str, Any]]] = []
    for raw_year, summary in yearly_summary.items():
        try:
            year = int(raw_year)
        except (TypeError, ValueError):
            continue
        output.append((year, summary))
    return sorted(output, key=lambda item: item[0])


def _sheet_has_data(sheet: Worksheet, header_row: int, first_row: int, last_row: int) -> bool:
    headers = table_headers(sheet, header_row)
    if not headers:
        return False

    for row_idx in range(first_row, last_row + 1):
        for col_idx in headers.values():
            if sheet.cell(row=row_idx, column=col_idx).value not in (None, ""):
                return True
    return False


def _require_capacity(label: str, count: int, first_row: int, last_row: int) -> None:
    capacity = last_row - first_row + 1
    if count > capacity:
        raise ExcelWriterError(
            f"{label} has room for {capacity} rows, but the engine needs to write {count}."
        )


def _zero_if_blank(value: Any) -> Any:
    return 0 if value in (None, "") else value


def _rows_checked(item: dict[str, Any]) -> Any:
    return item.get("rows_checked") or 0


def _checked_count(item: dict[str, Any], field: str) -> Any:
    return item.get(field) if _rows_checked(item) else "Not Checked"


def _catalog_notes(item: dict[str, Any]) -> str:
    if not item.get("exists", True):
        return "File missing."
    if item.get("status") == "Pending":
        return "Discovered by engine; run with price validation for row/date checks."
    return f"{item.get('report_status', item.get('status'))}; {_rows_checked(item)} price rows checked in this run."


def _quality_notes(item: dict[str, Any]) -> str:
    notes: list[str] = []
    if not item.get("exists", True):
        notes.append("File missing")
    missing = item.get("missing_required_columns") or []
    if missing:
        notes.append("Missing columns: " + ", ".join(str(value) for value in missing))
    if item.get("duplicate_timestamps"):
        notes.append(f"Duplicate timestamps: {item.get('duplicate_timestamps')}")
    if item.get("out_of_order_timestamps"):
        notes.append(f"Out-of-order timestamps: {item.get('out_of_order_timestamps')}")
    if item.get("bad_bid_ask_rows"):
        notes.append(f"Bad bid/ask rows: {item.get('bad_bid_ask_rows')}")
    if item.get("negative_spread_rows"):
        notes.append(f"Negative spread rows: {item.get('negative_spread_rows')}")
    if not item.get("validation_performed"):
        notes.append("Not validated in this run")
    elif not item.get("rows_checked"):
        notes.append("Quick validation only; full row scan not requested")
    elif item.get("rows_checked", 0) < (item.get("rows") or 0):
        notes.append(f"Partial row scan: {item['rows_checked']} of {item['rows']} rows")
    else:
        notes.append(f"Full row scan: {item['rows_checked']} rows")
    if item.get("date_range_ok") is False:
        notes.append("Empty or invalid timestamp bounds/year/order")
    if item.get("row_check_scope"):
        notes.append(item["row_check_scope"])
    if item.get("date_range_scope"):
        notes.append(item["date_range_scope"])
    diagnostics = item.get("diagnostics") or {}
    if diagnostics:
        notes.append("Additional diagnostics are informational, not trade filters")
        notes.extend(diagnostics.get("warnings", []))
        notes.append("Max_Spread is in pips, from candle open/close samples, not the tick peak")
        notes.append("Missing_Bars uses a reference FX schedule; unverified holidays/no quotes may explain gaps")
        notes.append("Timezone Consistent is not proof of source timezone; Review does not block execution")
        notes.append("See archived diagnostics for gap times, max-spread candle and timezone evidence")
    else:
        notes.append("Missing-bar, spread and timezone diagnostics require a row scan")
    return "; ".join(notes)


def normalize_result(outcome: Any) -> str:
    if outcome == "Win":
        return "Win"
    if outcome == "Loss":
        return "Loss"
    if outcome == "Breakeven":
        return "BE"
    if outcome == "No fill":
        return "No Fill"
    if outcome == "Blocked":
        return "Blocked"
    return "Open"


def relative_to_root(path: Path, root: Path) -> str:
    try:
        return str(path.resolve().relative_to(root.resolve()))
    except ValueError:
        return str(path)
