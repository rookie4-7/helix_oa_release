"""Runner reporting tests; no strategy rules or historical backtests are exercised."""
import argparse
from datetime import datetime, timedelta
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook

# Keep reporting tests independent of the selected strategy implementation.
with patch.dict(sys.modules, {"strategy_logic": SimpleNamespace(
    run_strategy_for_period=Mock(), strategy_config=Mock(),
)}):
    import backtest_engine
    import data_quality as quality
    import excel_runner
    import excel_writer
    from models import PriceFile, SelectedWindow
    from price_data_loader import inspect_price_file, stats_to_manifest_item


CATALOG_HEADERS = [
    "Data_ID", "Symbol", "Timeframe", "Data_Type", "File_Path", "File_Format",
    "Start_Date", "End_Date", "Rows", "Timezone", "Data_Version", "Checksum",
    "Status", "Last_Validated", "Quality_Report_ID", "Notes",
]
REPORT_HEADERS = [
    "Report_ID", "Data_ID", "Check_DateTime", "Check_Status", "Missing_Bars",
    "Duplicate_Timestamps", "Bad_Bid_Ask", "Negative_Spread", "Max_Spread",
    "Timezone_OK", "Date_Range_OK", "Rows_Checked", "Engine_Version", "Report_Path", "Notes",
]


class DataQualityReporting(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.file = PriceFile("GBP_USD", "1m", 2022, self.root / "prices.xlsx", "2022", "yearly_excel")
        self.headers = list(quality.REQUIRED_MBA_COLUMNS)
        self.now = datetime(2022, 1, 3, 8)

    def price_row(self, timestamp):
        return [timestamp if h == "time" else 1.2 if h.startswith("mid_")
                else 1.1999 if h.startswith("bid_") else 1.2001 for h in self.headers]

    def fixture(self, rows=None, headers=None):
        book = Workbook()
        sheet = book.active
        sheet.title = "2022"
        sheet.append(headers if headers is not None else self.headers)
        for row in (rows if rows is not None else [self.price_row(self.now), self.price_row(self.now + timedelta(minutes=1))]):
            sheet.append(row)
        book.save(self.file.path)
        book.close()
        return self.file

    def item(self, scan=True, limit=None):
        stats = inspect_price_file(self.file, scan_rows=scan, max_scan_rows=limit)
        return stats_to_manifest_item(stats, self.root)

    def table_book(self):
        book = Workbook()
        catalog = book.active
        catalog.title = "Data_Catalog"
        report = book.create_sheet("Data_Quality_Report")
        for sheet, headers in ((catalog, CATALOG_HEADERS), (report, REPORT_HEADERS)):
            for col, header in enumerate(headers, 1):
                sheet.cell(3, col, header)
        return book

    def report(self, item):
        book = self.table_book()
        self.addCleanup(book.close)
        path = self.root / "archive/mark_4_current_engine/RUN_0042/data_manifest.json"
        excel_writer.write_data_catalog(book["Data_Catalog"], {"files": [item]}, self.now)
        excel_writer.write_data_quality_report(book["Data_Quality_Report"], {"files": [item]},
                                               "RUN_0042", "Mark4-0.1.0", self.root, self.now,
                                               manifest_path=path)
        return (dict(zip(CATALOG_HEADERS, [c.value for c in book["Data_Catalog"][4]])),
                dict(zip(REPORT_HEADERS, [c.value for c in book["Data_Quality_Report"][4]])))

    def test_full_scan_reports_validated_counts_dates_and_actual_archive(self):
        self.fixture()
        catalog, report = self.report(self.item())
        self.assertEqual(catalog["Status"], "Review")
        self.assertEqual(catalog["Last_Validated"], self.now)
        self.assertEqual(report["Check_Status"], "Review")
        self.assertEqual(report["Rows_Checked"], 2)
        self.assertEqual(report["Duplicate_Timestamps"], 0)
        self.assertEqual(report["Date_Range_OK"], "Yes")
        self.assertEqual(report["Report_Path"], "archive/mark_4_current_engine/RUN_0042/data_manifest.json")
        self.assertIn("Full row scan: 2", report["Notes"])

    def test_additional_checks_are_reported_without_changing_core_status(self):
        self.fixture()
        item = self.item()
        catalog, report = self.report(item)
        self.assertEqual(catalog["Timezone"], "UTC")
        self.assertEqual(item["status"], "Validated")
        self.assertEqual(item["report_status"], "Review")
        self.assertGreater(report["Missing_Bars"], 0)
        self.assertEqual(report["Max_Spread"], 2.0)
        self.assertEqual(report["Timezone_OK"], "Inconclusive")
        self.assertIn("not calendar completeness", report["Notes"])
        self.assertIn("informational, not trade filters", report["Notes"])

    def test_quick_scan_is_not_full_validation_or_all_rows_checked(self):
        self.fixture()
        _, report = self.report(self.item(scan=False))
        self.assertEqual(report["Check_Status"], "Headers Checked")
        self.assertEqual(report["Rows_Checked"], 0)
        self.assertEqual(report["Bad_Bid_Ask"], "Not Checked")
        self.assertEqual(report["Date_Range_OK"], "Yes")
        for key in ("Missing_Bars", "Max_Spread", "Timezone_OK"):
            self.assertEqual(report[key], "Not Checked")

    def test_partial_scan_is_labelled_partial(self):
        self.fixture()
        _, report = self.report(self.item(limit=1))
        self.assertEqual(report["Check_Status"], "Partially Checked")
        self.assertEqual(report["Rows_Checked"], 1)
        self.assertIn("Partial row scan: 1 of 2", report["Notes"])
        self.assertEqual(report["Missing_Bars"], "Partial scan")
        self.assertEqual(report["Timezone_OK"], "Partial scan")

    def test_pending_counts_and_cached_dates_are_not_reported_as_checked(self):
        _, report = self.report({"data_id": "OLD", "status": "Pending", "rows": 100,
                                 "start_date": self.now, "end_date": self.now, "timezone": "UTC"})
        self.assertEqual(report["Rows_Checked"], 0)
        self.assertIsNone(report["Check_DateTime"])
        self.assertEqual(report["Duplicate_Timestamps"], "Not Checked")
        self.assertEqual(report["Date_Range_OK"], "Not Checked")

    def test_report_units_and_status_choices(self):
        self.fixture()
        book = self.table_book()
        self.addCleanup(book.close)
        sheet = book["Data_Quality_Report"]
        for _ in range(2):
            excel_writer.write_data_quality_report(sheet, {"files": [self.item()]},
                "RUN_1", "Mark4", self.root, self.now)
        self.assertEqual(sheet["I4"].number_format, '0.00" pips"')
        self.assertIsNone(sheet["J3"].comment)
        validations = {str(v.sqref): v.formula1 for v in sheet.data_validations.dataValidation}
        self.assertEqual(len(validations), 3)
        self.assertIn("Consistent", validations["J4:J35"])
        self.assertIn("Review", validations["D4:D35"])

    def test_empty_file_is_not_validated(self):
        self.fixture(rows=[])
        self.assertEqual(self.item()["status"], "Failed")

    def test_missing_file_is_not_validated(self):
        self.assertEqual(self.item()["status"], "Failed")

    def test_missing_required_columns_are_reported(self):
        self.fixture(rows=[], headers=["time"])
        item = self.item()
        self.assertEqual(item["status"], "Failed")
        self.assertIn("bid_close", item["missing_required_columns"])

    def test_wrong_year_bounds_fail(self):
        self.fixture(rows=[self.price_row(datetime(2021, 1, 3))])
        self.assertFalse(self.item()["date_range_ok"])
        self.assertEqual(self.item()["status"], "Failed")

    def test_duplicate_timestamps_fail(self):
        self.fixture(rows=[self.price_row(self.now), self.price_row(self.now)])
        self.assertEqual(self.item()["duplicate_timestamps"], 1)
        self.assertEqual(self.item()["status"], "Failed")

    def test_out_of_order_timestamps_fail(self):
        self.fixture(rows=[self.price_row(self.now), self.price_row(self.now - timedelta(minutes=1))])
        self.assertEqual(self.item()["out_of_order_timestamps"], 1)
        self.assertFalse(self.item()["date_range_ok"])

    def test_negative_spread_fails(self):
        row = self.price_row(self.now)
        row[self.headers.index("bid_close")] = 1.2002
        self.fixture(rows=[row])
        item = self.item()
        self.assertEqual(item["negative_spread_rows"], 1)
        self.assertEqual(item["status"], "Failed")

    def test_missing_close_quote_fails(self):
        row = self.price_row(self.now)
        row[self.headers.index("ask_close")] = None
        self.fixture(rows=[row])
        self.assertEqual(self.item()["bad_bid_ask_rows"], 1)

    def test_nonfinite_quotes_are_not_numbers(self):
        for value in ("nan", "inf", "-inf", "abc", None):
            with self.subTest(value=value):
                self.assertIsNone(quality.as_number(value))

    def test_blank_header_does_not_shift_quote_columns(self):
        headers = ["", *self.headers]
        rows = [[None, *self.price_row(self.now)]]
        self.fixture(rows=rows, headers=headers)
        self.assertEqual(self.item()["status"], "Validated")
        self.assertEqual(self.item(scan=False)["start_date"], self.now)

    def test_manifest_refreshes_catalog_and_uses_only_selected_inputs(self):
        self.fixture()
        book = self.table_book()
        self.addCleanup(book.close)
        # A previous run's stale/missing file must not substitute for this run's inputs.
        book["Data_Catalog"].cell(4, 1, "GBP_USD_1m_2021")
        book["Data_Catalog"].cell(4, 5, "missing_2021.xlsx")
        selected = SelectedWindow("WF_2022", datetime(2020, 1, 1), datetime(2021, 12, 31),
                                  datetime(2022, 1, 1), datetime(2022, 12, 31), 2022, "Mark4")
        for cached_status in ("Pending", "Failed", "Validated"):
            with self.subTest(status=cached_status):
                book["Data_Catalog"].cell(4, 13, cached_status)
                with patch.object(backtest_engine, "discover_price_files", return_value=[self.file]) as discovery:
                    result = backtest_engine.read_data_manifest(book, self.root / "input.xlsm", self.root,
                        self.root, selected, ["OOS"], False, True, True, None)
                discovery.assert_called_once_with(self.root, years=[2022])
                self.assertEqual(result["source"], "validated_selected_price_files")
                self.assertEqual(result["files"][0]["data_id"], "GBP_USD_1m_2022")
                self.assertEqual(result["files"][0]["status"], "Validated")
                self.assertEqual(result["files"][0]["rows_checked"], 2)

    def test_hash_requested_is_computed_for_discovered_files(self):
        self.fixture()
        book = self.table_book()
        self.addCleanup(book.close)
        selected = SelectedWindow("WF_2022", None, None, self.now, self.now, 2022, "Mark4")
        with patch.object(backtest_engine, "discover_price_files", return_value=[self.file]):
            result = backtest_engine.read_data_manifest(book, self.root / "input.xlsm", self.root,
                self.root, selected, ["OOS"], True, True, True, None)
        self.assertEqual(result["files"][0]["sha256"], backtest_engine.sha256_file(self.file.path))

    def test_excel_button_enables_full_scan_in_both_modes(self):
        for mode in ("dry-run", "write"):
            args = argparse.Namespace(workbook=str(self.root / "input.xlsm"), mode=mode,
                project_root=str(self.root), output_root=str(self.root), output_workbook=None,
                activate=True, official_oos=False, disable_news_filter=False, disable_news_forced_close=False)
            result = excel_runner.build_engine_args(args)
            self.assertTrue(result.validate_price_data)
            self.assertTrue(result.scan_price_data)
            self.assertIsNone(result.max_scan_rows)


if __name__ == "__main__":
    unittest.main()
